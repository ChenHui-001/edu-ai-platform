"""知识点附件文本提取（供 RAG 索引使用）。

不引入新的重量级依赖：
- PDF 用已安装的 PyMuPDF（fitz）；
- xlsx 用 openpyxl；
- docx/pptx 本质是 zip+XML，直接解析文本节点；
- txt/md/csv 按 UTF-8/GB18030 解码。
图片/旧版 Office（doc/xls/ppt）暂不解析，返回空字符串，不影响附件上传与下载。
"""
from __future__ import annotations

import logging
import os
import zipfile
from typing import List
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)

MAX_CHARS = 60000
MAX_PDF_PAGES = 40
MAX_SHEET_ROWS = 500
TEXT_EXTS = {".txt", ".md", ".csv"}


def _truncate(text: str) -> str:
    text = (text or "").strip()
    return text[:MAX_CHARS]


def _pdf_text(path: str) -> str:
    try:
        import fitz  # PyMuPDF
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF 解析不可用: %s", exc)
        return ""
    parts: List[str] = []
    try:
        with fitz.open(path) as doc:
            for page in doc[:MAX_PDF_PAGES]:
                parts.append(page.get_text("text") or "")
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF 文本提取失败 %s: %s", path, exc)
        return ""
    return _truncate("\n".join(parts))


def _xlsx_text(path: str) -> str:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # noqa: BLE001
        logger.warning("xlsx 解析不可用: %s", exc)
        return ""
    parts: List[str] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            parts.append(f"工作表：{ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    break
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
    except Exception as exc:  # noqa: BLE001
        logger.warning("xlsx 文本提取失败 %s: %s", path, exc)
        return ""
    return _truncate("\n".join(parts))


# 解压体积上限。docx/pptx 本质是 zip，压缩比可达 1000:1 —— 一个 1MB 的
# 文件能解出 1GB。附件本身受 MAX_UPLOAD_SIZE（50MB）约束，但那只管
# **压缩后**大小，挡不住 zip bomb。提取阶段必须按解压后体积再设一道闸，
# 否则索引/RAG 的 BackgroundTask 会把单进程容器直接 OOM（整站不可用）。
_MAX_XML_ENTRY_BYTES = 8 * 1024 * 1024    # 单个 XML 条目
_MAX_XML_TOTAL_BYTES = 32 * 1024 * 1024   # 整包累计
_MAX_PLAIN_BYTES = 4 * 1024 * 1024        # 纯文本类附件


def _zip_xml_text(path: str, prefix: str, tag: str) -> str:
    """从 docx/pptx 的 zip 包里提取指定 XML 标签文本。"""
    texts: List[str] = []
    total = 0
    try:
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist()
                     if n.startswith(prefix) and n.endswith(".xml") and "/_rels/" not in n]
            for name in sorted(names):
                try:
                    info = zf.getinfo(name)
                except KeyError:
                    continue
                # 必须用 ZipInfo.file_size 在 read() **之前**判断：
                # 读完再判断，内存已经被撑爆了。
                if info.file_size > _MAX_XML_ENTRY_BYTES:
                    logger.warning(
                        "跳过超大 XML 条目 %s（解压后 %d bytes，上限 %d）",
                        name, info.file_size, _MAX_XML_ENTRY_BYTES)
                    continue
                total += info.file_size
                if total > _MAX_XML_TOTAL_BYTES:
                    logger.warning(
                        "Office XML 累计解压体积超限（%d bytes），停止提取：%s",
                        total, path)
                    break
                try:
                    root = ET.fromstring(zf.read(name))
                except ET.ParseError:
                    continue
                for node in root.iter():
                    if node.tag.endswith("}" + tag) or node.tag == tag:
                        if node.text and node.text.strip():
                            texts.append(node.text.strip())
    except Exception as exc:  # noqa: BLE001
        logger.warning("Office XML 文本提取失败 %s: %s", path, exc)
        return ""
    return _truncate("\n".join(texts))


def _plain_text(path: str) -> str:
    # 超大纯文本同样会整份进内存（上传侧限的是 50MB，仍够撑爆小内存容器），
    # 提取阶段只处理前 4MB —— 索引/RAG 用不到更多。
    try:
        if os.path.getsize(path) > _MAX_PLAIN_BYTES:
            logger.warning(
                "纯文本附件超过提取上限（%d bytes），只读取前 %d bytes：%s",
                os.path.getsize(path), _MAX_PLAIN_BYTES, path)
            with open(path, "rb") as f:
                raw = f.read(_MAX_PLAIN_BYTES)
        else:
            raw = open(path, "rb").read()
    except OSError as exc:
        logger.warning("读取附件失败 %s: %s", path, exc)
        return ""
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return _truncate(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    return _truncate(raw.decode("utf-8", errors="ignore"))


def extract_text(file_path: str, filename: str = "") -> str:
    """提取附件文本；不支持/失败返回空字符串（调用方应视为不可索引）。"""
    if not file_path or not os.path.exists(file_path):
        return ""
    ext = os.path.splitext(filename or file_path)[1].lower()
    if ext == ".pdf":
        return _pdf_text(file_path)
    if ext == ".xlsx":
        return _xlsx_text(file_path)
    if ext == ".docx":
        return _zip_xml_text(file_path, "word/", "t")
    if ext == ".pptx":
        return _zip_xml_text(file_path, "ppt/slides/", "t")
    if ext in TEXT_EXTS:
        return _plain_text(file_path)
    return ""
