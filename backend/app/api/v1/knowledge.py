"""知识点管理 V2：树形结构、数据范围权限、批量操作、导入导出。

契约见 docs/KNOWLEDGE_V2.md（阶段 1+2）。要点：
- 树由服务端维护 level/path，不接受客户端伪造；
- 读取按身份收敛（学生强制本年级 + published），写入按 scope 校验（越权 403）；
- 旧接口 GET /edu/knowledge、/edu/knowledge/list 保持「平铺数组」兼容 E2E；
- POST /edu/knowledge/reindex 保留教师可调用（E2E 既有契约）。
"""
import csv
import hashlib
import io
import os
import time
import uuid
from typing import Any, Dict, List, Optional, Set, Tuple
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, File, Form, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import case, func, or_, select
from sqlalchemy.orm import Session

from app.agents import rag as rag_engine
from app.agents.rag import (
    delete_knowledge_index, index_knowledge, knowledge_index_stats, reindex_all,
)
from app.api.deps import (
    DataScope, get_current_teacher, get_current_user, get_data_scope, require_perm,
)
from app.api.v1.attachment import OFFICE_TYPES, PREVIEW_TYPES
from app.core.config import settings
from app.core.database import get_db
from app.core.exception import BusinessException, NotFoundException, PermissionException
from app.models import (
    Attachment, Class, Grade, KnowledgeMastery, KnowledgePoint, Question, ReviewPlan,
    Subject, TeachingAssignment, User,
)
from app.utils.common import dictify, now_local, ok

router = APIRouter(prefix="/edu", tags=["知识点"])

VALID_STATUS = {"draft", "published", "archived"}
VALID_DIFFICULTY = {"easy", "medium", "hard"}
VALID_SOURCE = {"manual", "import", "ai", "seed"}
STATUS_ALL = "all"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IMPORT_HEADERS = [
    "年级", "学科", "教材版本", "学期", "单元", "章节", "知识点", "难度", "关键词", "说明", "编码",
]
WRITE_PERM = "resource:knowledge:write"


# ================================================================ 权限 scope
def _codes(user: User) -> Set[str]:
    return set(user.role_codes or [])


def _is_sys_admin(user: User) -> bool:
    return "sys_admin" in _codes(user)


def _is_admin(user: User) -> bool:
    """系统管理员或年级管理员（可查看草稿/归档节点）。"""
    return bool({"sys_admin", "grade_admin"} & _codes(user))


def _is_student_user(user: User) -> bool:
    """纯学生视角（无任何教师/管理角色）。"""
    return user.identity == "student" and not (
        {"sys_admin", "grade_admin", "teacher"} & _codes(user))


def _teacher_pairs(db: Session, user: User) -> Set[Tuple[int, int]]:
    rows = db.execute(
        select(TeachingAssignment.grade_id, TeachingAssignment.subject_id).where(
            TeachingAssignment.teacher_id == user.id,
            TeachingAssignment.is_active.is_(True),
        )
    ).all()
    return {(int(g), int(s)) for g, s in rows if g and s}


def _write_predicate(db: Session, user: User):
    """返回 (grade_id, subject_id) -> bool 的写权限判定函数。

    批量/导入场景会反复判定，这里一次性取出任教关系，避免逐节点查库。
    """
    codes = _codes(user)
    if "sys_admin" in codes:
        return lambda grade_id, subject_id: True
    if "grade_admin" in codes:
        gid = user.grade_admin_grade_id or user.grade_id
        return lambda grade_id, subject_id: gid is not None and int(gid) == int(grade_id)
    if "teacher" in codes:
        pairs = _teacher_pairs(db, user)
        return lambda grade_id, subject_id: (int(grade_id), int(subject_id)) in pairs
    return lambda grade_id, subject_id: False


def _can_write(db: Session, user: User, grade_id: int, subject_id: int) -> bool:
    return _write_predicate(db, user)(grade_id, subject_id)


def _assert_write(db: Session, user: User, grade_id: int, subject_id: int) -> None:
    """写权限校验：越权统一 403（不暴露资源是否存在）。"""
    if not _can_write(db, user, grade_id, subject_id):
        raise PermissionException("无权维护该年级/学科的知识点")


def _effective_status(user: User, status: Optional[str]) -> Optional[str]:
    """读取状态过滤。

    - 默认只读 published；
    - status=all / draft / archived 仅系统/年级管理员可见，其他角色按 published 处理；
    - 非法值 400。
    """
    value = (status or "published").strip().lower()
    if value not in {"published", "draft", "archived", STATUS_ALL}:
        raise BusinessException("status 仅支持 published / draft / archived / all")
    if value == STATUS_ALL:
        return None if _is_admin(user) else "published"
    if value in {"draft", "archived"} and not _is_admin(user):
        return "published"
    return value


def _forced_grade(user: User, grade_id: Optional[int]) -> Optional[int]:
    """学生读写都强制本年级；显式请求其他年级直接 403。"""
    if not _is_student_user(user):
        return grade_id
    own = user.grade_id
    if grade_id and own and int(grade_id) != int(own):
        raise PermissionException("学生只能查看本年级的知识点")
    return int(own) if own else -1  # -1 表示无年级 → 查不到数据而不是放行


# ================================================================ 通用工具
def _to_int(value: Any, default: Optional[int] = None) -> Optional[int]:
    if value in (None, ""):
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _clean_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _validate_name(name: Any) -> str:
    text = str(name or "").strip()
    if not text:
        raise BusinessException("知识点名称不能为空")
    if len(text) > 128:
        raise BusinessException("知识点名称不能超过 128 个字符")
    return text


def _validate_difficulty(value: Any) -> str:
    diff = str(value or "medium").strip().lower()
    if diff not in VALID_DIFFICULTY:
        raise BusinessException("难度只能是 easy / medium / hard")
    return diff


def _validate_status(value: Any, default: str = "published") -> str:
    status = str(value or default).strip().lower()
    if status not in VALID_STATUS:
        raise BusinessException("状态只能是 draft / published / archived")
    return status


def _sibling_exists(
    db: Session, grade_id: int, subject_id: int, parent_id: Optional[int],
    name: str, exclude_id: Optional[int] = None,
) -> bool:
    stmt = select(KnowledgePoint.id).where(
        KnowledgePoint.grade_id == grade_id,
        KnowledgePoint.subject_id == subject_id,
        KnowledgePoint.name == name,
    )
    stmt = stmt.where(KnowledgePoint.parent_id.is_(None)) if parent_id is None \
        else stmt.where(KnowledgePoint.parent_id == parent_id)
    if exclude_id:
        stmt = stmt.where(KnowledgePoint.id != exclude_id)
    return db.scalar(stmt.limit(1)) is not None


def _code_exists(
    db: Session, grade_id: int, subject_id: int, code: Optional[str],
    exclude_id: Optional[int] = None,
) -> bool:
    if not code:
        return False
    stmt = select(KnowledgePoint.id).where(
        KnowledgePoint.grade_id == grade_id,
        KnowledgePoint.subject_id == subject_id,
        KnowledgePoint.code == code,
    )
    if exclude_id:
        stmt = stmt.where(KnowledgePoint.id != exclude_id)
    return db.scalar(stmt.limit(1)) is not None


def _children_map(db: Session, grade_id: int, subject_id: int) -> Dict[int, List[int]]:
    rows = db.execute(
        select(KnowledgePoint.id, KnowledgePoint.parent_id).where(
            KnowledgePoint.grade_id == grade_id,
            KnowledgePoint.subject_id == subject_id,
        )
    ).all()
    mapping: Dict[int, List[int]] = {}
    for cid, pid in rows:
        if pid:
            mapping.setdefault(pid, []).append(cid)
    return mapping


def _subtree_nodes(db: Session, kp: KnowledgePoint) -> List[KnowledgePoint]:
    """取节点全部后代（父先于子）。"""
    mapping = _children_map(db, kp.grade_id, kp.subject_id)
    out: List[KnowledgePoint] = []
    stack = [kp]
    while stack:
        cur = stack.pop()
        for cid in mapping.get(cur.id, []):
            child = db.get(KnowledgePoint, cid)
            if child:
                out.append(child)
                stack.append(child)
    return out


def _refresh_level_path(db: Session, kp: KnowledgePoint) -> None:
    """重算自身与全部后代的 level/path（parent 变更后调用）。"""
    parent = db.get(KnowledgePoint, kp.parent_id) if kp.parent_id else None
    if parent:
        kp.level = (parent.level or 0) + 1
        kp.path = f"{parent.path or parent.id}/{kp.id}"
    else:
        kp.level = 1
        kp.path = str(kp.id)
    nodes: Dict[int, KnowledgePoint] = {kp.id: kp}
    for child in _subtree_nodes(db, kp):
        nodes[child.id] = child
    for child in _subtree_nodes(db, kp):
        p = nodes.get(child.parent_id)
        if p:
            child.level = (p.level or 0) + 1
            child.path = f"{p.path or p.id}/{child.id}"


# ================================================================ 查询/序列化
def _filtered_query(
    db: Session, user: User, *, grade_id: Optional[int] = None,
    subject_id: Optional[int] = None, term: Optional[int] = None,
    keyword: Optional[str] = None, status: Optional[str] = None,
    subject_name: Optional[str] = None, level: Optional[int] = None,
):
    """按读取契约构造知识点查询（含身份收敛）。"""
    grade_id = _forced_grade(user, grade_id)
    effective_status = _effective_status(user, status)
    stmt = select(KnowledgePoint).where(KnowledgePoint.is_active.is_(True))
    if effective_status:
        stmt = stmt.where(KnowledgePoint.status == effective_status)
    if grade_id:
        stmt = stmt.where(KnowledgePoint.grade_id == grade_id)
    if subject_id:
        stmt = stmt.where(KnowledgePoint.subject_id == subject_id)
    if subject_name:
        stmt = stmt.where(KnowledgePoint.subject_id.in_(
            select(Subject.id).where(Subject.name == subject_name)
        ))
    if term:
        stmt = stmt.where(KnowledgePoint.term == term)
    if level:
        stmt = stmt.where(KnowledgePoint.level == level)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(
            KnowledgePoint.name.like(like),
            KnowledgePoint.keywords.like(like),
            KnowledgePoint.code.like(like),
        ))
    return stmt.order_by(KnowledgePoint.sort, KnowledgePoint.id)


def _name_maps(db: Session, rows: List[KnowledgePoint]) -> Tuple[Dict[int, str], Dict[int, str]]:
    grade_ids = {r.grade_id for r in rows if r.grade_id}
    subject_ids = {r.subject_id for r in rows if r.subject_id}
    grade_names = {
        g.id: g.name for g in db.scalars(
            select(Grade).where(Grade.id.in_(grade_ids))).all()
    } if grade_ids else {}
    subject_names = {
        s.id: s.name for s in db.scalars(
            select(Subject).where(Subject.id.in_(subject_ids))).all()
    } if subject_ids else {}
    return grade_names, subject_names


def _question_counts(db: Session, ids: List[int]) -> Dict[int, int]:
    if not ids:
        return {}
    rows = db.execute(
        select(Question.knowledge_point_id, func.count()).where(
            Question.knowledge_point_id.in_(ids)
        ).group_by(Question.knowledge_point_id)
    ).all()
    return {int(kp_id): int(cnt or 0) for kp_id, cnt in rows if kp_id}


def _node(
    kp: KnowledgePoint, grade_names: Dict[int, str], subject_names: Dict[int, str],
    question_counts: Dict[int, int], child_counts: Optional[Dict[int, int]] = None,
    include_children: bool = False,
) -> Dict[str, Any]:
    data = dictify(kp)
    data["grade_name"] = grade_names.get(kp.grade_id)
    data["subject_name"] = subject_names.get(kp.subject_id)
    data["question_count"] = int(question_counts.get(kp.id, 0))
    data["child_count"] = int((child_counts or {}).get(kp.id, 0))
    if include_children:
        data["children"] = []
    return data


def _flat_nodes(db: Session, rows: List[KnowledgePoint]) -> List[Dict[str, Any]]:
    grade_names, subject_names = _name_maps(db, rows)
    question_counts = _question_counts(db, [r.id for r in rows])
    child_counts: Dict[int, int] = {}
    for r in rows:
        if r.parent_id:
            child_counts[r.parent_id] = child_counts.get(r.parent_id, 0) + 1
    return [
        _node(r, grade_names, subject_names, question_counts, child_counts)
        for r in rows
    ]


def _tree_nodes(db: Session, rows: List[KnowledgePoint]) -> List[Dict[str, Any]]:
    grade_names, subject_names = _name_maps(db, rows)
    question_counts = _question_counts(db, [r.id for r in rows])
    child_counts: Dict[int, int] = {}
    for r in rows:
        if r.parent_id:
            child_counts[r.parent_id] = child_counts.get(r.parent_id, 0) + 1
    nodes: Dict[int, Dict[str, Any]] = {}
    for r in rows:
        nodes[r.id] = _node(
            r, grade_names, subject_names, question_counts, child_counts,
            include_children=True,
        )
    roots: List[Dict[str, Any]] = []
    for r in rows:
        node = nodes[r.id]
        parent = nodes.get(r.parent_id) if r.parent_id else None
        if parent is not None:
            parent["children"].append(node)
        else:
            roots.append(node)
    return roots


def _single_node(db: Session, kp: KnowledgePoint) -> Dict[str, Any]:
    grade_names, subject_names = _name_maps(db, [kp])
    question_counts = _question_counts(db, [kp.id])
    child_count = db.scalar(
        select(func.count()).select_from(KnowledgePoint).where(
            KnowledgePoint.parent_id == kp.id)
    ) or 0
    return _node(kp, grade_names, subject_names, question_counts,
                 {kp.id: int(child_count)})


def _breadcrumb(db: Session, kp: KnowledgePoint) -> List[Dict[str, Any]]:
    chain: List[Dict[str, Any]] = []
    cur: Optional[KnowledgePoint] = kp
    guard = 0
    while cur and guard < 50:
        chain.append({"id": cur.id, "name": cur.name})
        cur = db.get(KnowledgePoint, cur.parent_id) if cur.parent_id else None
        guard += 1
    chain.reverse()
    return chain


# ================================================================ 读取接口
@router.get("/knowledge/tree", summary="知识点树")
def knowledge_tree(
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    subject_name: Optional[str] = None, term: Optional[int] = None,
    keyword: Optional[str] = None, status: Optional[str] = None,
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    stmt = _filtered_query(
        db, user, grade_id=grade_id, subject_id=subject_id, term=term,
        keyword=keyword, status=status, subject_name=subject_name,
    )
    rows = list(db.scalars(stmt).all())
    return ok({"items": _tree_nodes(db, rows), "total": len(rows)})


@router.get("/knowledge", summary="知识点平铺列表（兼容旧接口）")
def knowledge_flat(
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    subject_name: Optional[str] = None, term: Optional[int] = None,
    keyword: Optional[str] = None, status: Optional[str] = None,
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    """旧接口：返回平铺数组（无 children），E2E 与旧前端依赖此结构。"""
    stmt = _filtered_query(
        db, user, grade_id=grade_id, subject_id=subject_id, term=term,
        keyword=keyword, status=status, subject_name=subject_name,
    )
    return ok(_flat_nodes(db, list(db.scalars(stmt).all())))


@router.get("/knowledge/list", summary="知识点平铺列表（兼容旧接口）")
def knowledge_list(
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    level: Optional[int] = None, keyword: Optional[str] = None,
    status: Optional[str] = None, params_limit: int = Query(200, le=1000),
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    stmt = _filtered_query(
        db, user, grade_id=grade_id, subject_id=subject_id, level=level,
        keyword=keyword, status=status,
    )
    rows = list(db.scalars(stmt.limit(params_limit)).all())
    return ok(_flat_nodes(db, rows))


@router.get("/knowledge/stats", summary="知识点统计")
def knowledge_stats(
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    term: Optional[int] = None, keyword: Optional[str] = None,
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    """统计各状态数量与年级/学科分布。

    状态维度始终输出完整分布（管理员可见 draft/archived），因此这里不套用
    status 过滤；非管理员按读取契约只看得到 published。
    """
    stmt = select(KnowledgePoint).where(KnowledgePoint.is_active.is_(True))
    visible_grade = _forced_grade(user, grade_id)
    if visible_grade:
        stmt = stmt.where(KnowledgePoint.grade_id == visible_grade)
    if subject_id:
        stmt = stmt.where(KnowledgePoint.subject_id == subject_id)
    if term:
        stmt = stmt.where(KnowledgePoint.term == term)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(
            KnowledgePoint.name.like(like), KnowledgePoint.keywords.like(like),
            KnowledgePoint.code.like(like)))
    rows = list(db.scalars(stmt).all())
    if not _is_admin(user):
        rows = [r for r in rows if r.status == "published"]

    question_counts = _question_counts(db, [r.id for r in rows])
    grade_ids = {r.grade_id for r in rows if r.grade_id}
    subject_ids = {r.subject_id for r in rows if r.subject_id}
    grade_names = {
        g.id: g.name for g in db.scalars(
            select(Grade).where(Grade.id.in_(grade_ids))).all()
    } if grade_ids else {}
    subject_names = {
        s.id: s.name for s in db.scalars(
            select(Subject).where(Subject.id.in_(subject_ids))).all()
    } if subject_ids else {}

    by_grade: Dict[int, int] = {}
    by_subject: Dict[int, int] = {}
    for r in rows:
        by_grade[r.grade_id] = by_grade.get(r.grade_id, 0) + 1
        by_subject[r.subject_id] = by_subject.get(r.subject_id, 0) + 1

    return ok({
        "total": len(rows),
        "published": sum(1 for r in rows if r.status == "published"),
        "draft": sum(1 for r in rows if r.status == "draft"),
        "archived": sum(1 for r in rows if r.status == "archived"),
        "with_questions": sum(1 for r in rows if question_counts.get(r.id)),
        "by_grade": [
            {"grade_id": gid, "grade_name": grade_names.get(gid), "count": cnt}
            for gid, cnt in sorted(by_grade.items())
        ],
        "by_subject": [
            {"subject_id": sid, "subject_name": subject_names.get(sid), "count": cnt}
            for sid, cnt in sorted(by_subject.items())
        ],
    })


@router.get("/knowledge/template", summary="下载知识点导入模板")
def knowledge_template(_: User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入模板"
    ws.append(IMPORT_HEADERS)
    widths = [10, 10, 14, 8, 16, 16, 20, 10, 18, 24, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    return _xlsx_response(wb, "知识点导入模板.xlsx")


@router.get("/knowledge/export", summary="导出知识点")
def knowledge_export(
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    term: Optional[int] = None, keyword: Optional[str] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    """按树展平导出 xlsx（列与导入模板一致，名称带层级缩进）。"""
    stmt = _filtered_query(
        db, user, grade_id=grade_id, subject_id=subject_id, term=term,
        keyword=keyword, status=status,
    )
    roots = _tree_nodes(db, list(db.scalars(stmt).all()))

    flat: List[Dict[str, Any]] = []
    stack = list(reversed(roots))
    while stack:
        node = stack.pop()
        flat.append(node)
        stack.extend(reversed(node.get("children") or []))

    wb = Workbook()
    ws = wb.active
    ws.title = "知识点"
    ws.append(IMPORT_HEADERS)
    for node in flat:
        level = int(node.get("level") or 1)
        unit = node["name"] if level == 1 else ""
        chapter = node["name"] if level == 2 else ""
        point = ("  " * (level - 3)) + node["name"] if level >= 3 else ""
        ws.append([
            node.get("grade_name") or "", node.get("subject_name") or "",
            node.get("textbook_version") or "", node.get("term") or 1,
            unit, chapter, point, node.get("difficulty") or "medium",
            node.get("keywords") or "", node.get("description") or "",
            node.get("code") or "",
        ])
    return _xlsx_response(wb, "知识点导出.xlsx")


@router.get("/knowledge/mastery-overview", summary="班级知识点掌握概览（教师端）")
def knowledge_mastery_overview(
    class_id: int = Query(..., description="班级 ID"),
    db: Session = Depends(get_db), user: User = Depends(get_current_teacher),
    scope: DataScope = Depends(get_data_scope),
):
    """聚合班内学生在每个知识点上的掌握数据，薄弱优先排序。

    原实现在 edu.py，V2 迁移到本文件，路径与返回结构保持不变。
    """
    if not scope.is_super:
        allowed = scope.class_ids()
        if allowed is not None and class_id not in allowed:
            if "grade_admin" not in user.role_codes:
                raise BusinessException("无权查看该班级的掌握数据")

    k = db.get(Class, class_id)
    if not k:
        raise NotFoundException("班级不存在")

    stu_ids = db.scalars(select(User.id).where(
        User.class_id == class_id, User.identity == "student",
        User.status == "active",
    )).all()
    if not stu_ids:
        return ok({"items": [], "student_count": 0, "class_name": k.name})

    rows = db.execute(
        select(
            KnowledgeMastery.knowledge_point_id,
            func.avg(KnowledgeMastery.mastery),
            func.sum(KnowledgeMastery.total_count),
            func.count(KnowledgeMastery.student_id),
            func.sum(case((KnowledgeMastery.mastery < 60, 1), else_=0)),
        ).where(
            KnowledgeMastery.student_id.in_(stu_ids)
        ).group_by(KnowledgeMastery.knowledge_point_id)
    ).all()

    kp_ids = {r[0] for r in rows}
    kp_map = {x.id: x for x in db.scalars(
        select(KnowledgePoint).where(KnowledgePoint.id.in_(kp_ids))).all()} if kp_ids else {}

    items = []
    for kp_id_, avg_m, total, stu_cnt, weak in rows:
        kp = kp_map.get(kp_id_)
        if not kp:
            continue
        items.append({
            "knowledge_point_id": kp_id_,
            "name": kp.name,
            "subject_id": kp.subject_id,
            "avg_mastery": round(float(avg_m or 0), 1),
            "total_count": int(total or 0),
            "student_count": int(stu_cnt or 0),
            "weak_count": int(weak or 0),
        })
    items.sort(key=lambda x: x["avg_mastery"])  # 薄弱优先
    return ok({"items": items, "student_count": len(stu_ids), "class_name": k.name})


# ================================================================ 语义检索 / RAG 状态
@router.get("/knowledge/search", summary="知识点语义搜索（RAG）")
def knowledge_search(
    q: str = Query(..., min_length=1, max_length=200, description="搜索问题/关键词"),
    grade_id: Optional[int] = None, subject_id: Optional[int] = None,
    limit: int = Query(10, ge=1, le=50),
    db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    """RAG 语义检索知识点（只返回 knowledge 类型结果）。

    读取权限按 1.1.0 契约收敛：学生强制本年级 + published；其他身份只返回
    published（索引本身也只收 published，这里再按 DB 兜底过滤存量脏索引）。
    """
    query = (q or "").strip()
    if not query:
        raise BusinessException("请输入搜索内容")

    student = _is_student_user(user)
    visible_grade = _forced_grade(user, grade_id)
    results = rag_engine.search(
        query, db, top_k=limit, grade_id=visible_grade,
        subject_id=subject_id, types=["knowledge"],
    )

    ids: List[int] = []
    for r in results:
        meta = r.get("metadata") or {}
        if meta.get("type") == "knowledge":
            value = _to_int(meta.get("id"))
            if value and value not in ids:
                ids.append(value)
    kps = {
        kp.id: kp for kp in db.scalars(
            select(KnowledgePoint).where(KnowledgePoint.id.in_(ids))
        ).all()
    } if ids else {}
    counts = _question_counts(db, list(kps))

    items: List[Dict[str, Any]] = []
    for r in results:
        meta = r.get("metadata") or {}
        if meta.get("type") != "knowledge":
            continue
        kp = kps.get(_to_int(meta.get("id")) or 0)
        if not kp or not kp.is_active or kp.status != "published":
            continue
        if student and (not user.grade_id or kp.grade_id != user.grade_id):
            continue
        if subject_id and kp.subject_id != int(subject_id):
            continue
        items.append({
            "id": kp.id, "name": kp.name, "path": kp.path or str(kp.id),
            "description": kp.description or "", "difficulty": kp.difficulty,
            "score": r.get("score"), "grade_id": kp.grade_id,
            "subject_id": kp.subject_id,
            "question_count": int(counts.get(kp.id, 0)),
        })
        if len(items) >= limit:
            break
    return ok({"items": items, "total": len(items)})


@router.get("/knowledge/rag-status", summary="知识点 RAG 索引状态")
def knowledge_rag_status(
    db: Session = Depends(get_db), user: User = Depends(get_current_teacher),
):
    """索引覆盖率（admin/teacher 可读）：供管理页展示与重建后刷新。"""
    return ok(knowledge_index_stats(db))


# ================================================================ 详情
@router.get("/knowledge/{kp_id}", summary="知识点详情")
def knowledge_detail(
    kp_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp:
        raise NotFoundException("知识点不存在")
    if _is_student_user(user):
        if not user.grade_id or kp.grade_id != user.grade_id:
            raise PermissionException("学生只能查看本年级的知识点")
        if kp.status != "published":
            raise NotFoundException("知识点不存在")
    node = _single_node(db, kp)
    node["breadcrumb"] = _breadcrumb(db, kp)
    return ok(node)


# ================================================================ 知识点附件
KNOWLEDGE_ATTACH_CATEGORY = "knowledge"


def _attachment_out(a: Attachment) -> Dict[str, Any]:
    """附件输出：不暴露物理路径/存储名，只给可安全使用的预览/下载地址。"""
    ext = os.path.splitext(a.filename or "")[1].lower()
    d = dictify(a)
    for key in ("file_path", "stored_name"):
        d.pop(key, None)
    d["ext"] = ext
    d["previewable"] = ext in PREVIEW_TYPES or ext in OFFICE_TYPES
    d["preview_url"] = f"/api/v1/attachments/{a.id}/preview"
    d["download_url"] = f"/api/v1/attachments/{a.id}/download"
    return d


def _get_attachment(db: Session, kp_id: int, att_id: int) -> Attachment:
    att = db.get(Attachment, att_id)
    if (not att or att.biz_type != "knowledge_point" or att.biz_id != kp_id):
        raise NotFoundException("附件不存在")
    return att


@router.get("/knowledge/{kp_id}/attachments", summary="知识点附件列表")
def list_knowledge_attachments(
    kp_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp or not kp.is_active:
        raise NotFoundException("知识点不存在")
    if _is_student_user(user):
        if not user.grade_id or kp.grade_id != user.grade_id or kp.status != "published":
            raise PermissionException("无权查看该知识点的附件")
    elif not _is_admin(user) and kp.status != "published" and not _can_write(
            db, user, kp.grade_id, kp.subject_id):
        raise PermissionException("无权查看该知识点的附件")

    stmt = select(Attachment).where(
        Attachment.biz_type == "knowledge_point", Attachment.biz_id == kp_id,
        Attachment.is_temp.is_(False),
    )
    if _is_student_user(user):
        stmt = stmt.where(Attachment.visible_to_student.is_(True))
    rows = db.scalars(stmt.order_by(Attachment.created_at.desc(), Attachment.id.desc())).all()
    return ok([_attachment_out(a) for a in rows])


@router.post("/knowledge/{kp_id}/attachments", summary="上传知识点附件")
async def upload_knowledge_attachment(
    kp_id: int, file: UploadFile = File(...),
    visible_to_student: bool = Form(True),
    db: Session = Depends(get_db), user: User = Depends(require_perm(WRITE_PERM)),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp or not kp.is_active:
        raise NotFoundException("知识点不存在")
    _assert_write(db, user, kp.grade_id, kp.subject_id)

    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in settings.ALLOWED_UPLOAD_TYPES:
        raise BusinessException(
            f"不支持的文件类型 {ext or '未知'}，允许：{'、'.join(settings.ALLOWED_UPLOAD_TYPES)}")
    content = await file.read()
    if not content:
        raise BusinessException("文件为空")
    max_bytes = settings.MAX_UPLOAD_SIZE * 1024 * 1024
    if len(content) > max_bytes:
        raise BusinessException(f"文件大小超过限制（{settings.MAX_UPLOAD_SIZE}MB）")

    digest = hashlib.md5(content).hexdigest()[:12]
    stored_name = f"{uuid.uuid4().hex}_{digest}{ext}"
    dir_path = os.path.join(settings.UPLOAD_DIR, now_local().strftime("%Y%m%d"))
    os.makedirs(dir_path, exist_ok=True)
    full_path = os.path.abspath(os.path.join(dir_path, stored_name))
    if not full_path.startswith(os.path.abspath(settings.UPLOAD_DIR)):
        raise BusinessException("非法文件路径")
    with open(full_path, "wb") as f:
        f.write(content)

    att = Attachment(
        filename=os.path.basename(file.filename or f"attachment{ext}")[:255],
        stored_name=stored_name, file_path=full_path,
        file_size=len(content), file_type=file.content_type,
        category=KNOWLEDGE_ATTACH_CATEGORY, biz_type="knowledge_point", biz_id=kp.id,
        uploader_id=user.id, is_temp=False, visible_to_student=bool(visible_to_student),
    )
    db.add(att)
    db.commit()
    db.refresh(att)
    return ok(_attachment_out(att), message="附件上传成功")


@router.put("/knowledge/{kp_id}/attachments/{att_id}", summary="更新知识点附件可见性")
def update_knowledge_attachment(
    kp_id: int, att_id: int, data: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db), user: User = Depends(require_perm(WRITE_PERM)),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp:
        raise NotFoundException("知识点不存在")
    _assert_write(db, user, kp.grade_id, kp.subject_id)
    att = _get_attachment(db, kp_id, att_id)
    if "visible_to_student" in data:
        att.visible_to_student = bool(data["visible_to_student"])
    db.commit()
    db.refresh(att)
    return ok(_attachment_out(att), message="附件已更新")


@router.delete("/knowledge/{kp_id}/attachments/{att_id}", summary="删除知识点附件")
def delete_knowledge_attachment(
    kp_id: int, att_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp:
        raise NotFoundException("知识点不存在")
    _assert_write(db, user, kp.grade_id, kp.subject_id)
    att = _get_attachment(db, kp_id, att_id)
    try:
        if os.path.exists(att.file_path):
            os.remove(att.file_path)
    except OSError:
        pass
    db.delete(att)
    db.commit()
    return ok(message="附件已删除")


# ================================================================ 写入接口
@router.post("/knowledge", summary="新增知识点")
def create_knowledge(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    name = _validate_name(data.get("name"))
    grade_id = _to_int(data.get("grade_id"))
    subject_id = _to_int(data.get("subject_id"))
    if not grade_id or not subject_id:
        raise BusinessException("请填写年级和学科")
    grade = db.get(Grade, grade_id)
    if not grade:
        raise BusinessException("年级不存在")
    subject = db.get(Subject, subject_id)
    if not subject:
        raise BusinessException("学科不存在")
    if subject.grade_id != grade_id:
        raise BusinessException("学科与年级不匹配")
    _assert_write(db, user, grade_id, subject_id)

    parent_id = _to_int(data.get("parent_id"))
    parent: Optional[KnowledgePoint] = None
    if parent_id:
        parent = db.get(KnowledgePoint, parent_id)
        if not parent:
            raise BusinessException("父级知识点不存在")
        if parent.grade_id != grade_id or parent.subject_id != subject_id:
            raise BusinessException("父级知识点必须与当前知识点同年级、同学科")
        # parent 与目标同属一个 (grade, subject)，写权限已在上方校验

    if _sibling_exists(db, grade_id, subject_id, parent_id, name):
        raise BusinessException("同级已存在同名知识点")
    code = _clean_str(data.get("code"))
    if _code_exists(db, grade_id, subject_id, code):
        raise BusinessException("该编码已存在")

    status = _validate_status(data.get("status"))
    source = str(data.get("source") or "manual").strip().lower()
    if source not in VALID_SOURCE:
        source = "manual"
    term = _to_int(data.get("term"), 1) or 1
    sort = _to_int(data.get("sort", data.get("sort_order")), 0) or 0
    kp = KnowledgePoint(
        name=name, grade_id=grade_id, subject_id=subject_id,
        parent_id=parent_id if parent else None, term=term, code=code,
        description=_clean_str(data.get("description")),
        difficulty=_validate_difficulty(data.get("difficulty")),
        keywords=_clean_str(data.get("keywords")), sort=sort,
        textbook_version=_clean_str(data.get("textbook_version")),
        status=status, source=source, owner_id=user.id, is_active=True,
    )
    db.add(kp)
    db.flush()
    # level/path 由服务端维护，忽略客户端伪造值
    if parent:
        kp.level = (parent.level or 0) + 1
        kp.path = f"{parent.path or parent.id}/{kp.id}"
    else:
        kp.level = 1
        kp.path = str(kp.id)
    db.commit()
    db.refresh(kp)
    index_knowledge(kp)  # 向量化失败不阻断创建
    return ok(_single_node(db, kp), message="知识点创建成功")


@router.put("/knowledge/{kp_id}", summary="编辑知识点")
def update_knowledge(
    kp_id: int, data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp:
        raise NotFoundException("知识点不存在")
    _assert_write(db, user, kp.grade_id, kp.subject_id)

    if "name" in data:
        name = _validate_name(data.get("name"))
        if name != kp.name and _sibling_exists(
                db, kp.grade_id, kp.subject_id, kp.parent_id, name, exclude_id=kp.id):
            raise BusinessException("同级已存在同名知识点")
        kp.name = name
    if "code" in data:
        code = _clean_str(data.get("code"))
        if code != kp.code and _code_exists(
                db, kp.grade_id, kp.subject_id, code, exclude_id=kp.id):
            raise BusinessException("该编码已存在")
        kp.code = code
    if "difficulty" in data:
        kp.difficulty = _validate_difficulty(data.get("difficulty"))
    if "status" in data:
        kp.status = _validate_status(data.get("status"), default=kp.status)
    if "term" in data:
        kp.term = _to_int(data.get("term"), kp.term) or kp.term
    if "sort" in data or "sort_order" in data:
        kp.sort = _to_int(data.get("sort", data.get("sort_order")), kp.sort) or 0
    for field in ("description", "keywords", "textbook_version"):
        if field in data:
            setattr(kp, field, _clean_str(data.get(field)))

    parent_changed = False
    if "parent_id" in data:
        new_parent_id = _to_int(data.get("parent_id"))
        if new_parent_id != kp.parent_id:
            parent_changed = True
            parent: Optional[KnowledgePoint] = None
            if new_parent_id:
                parent = db.get(KnowledgePoint, new_parent_id)
                if not parent:
                    raise BusinessException("父级知识点不存在")
                if parent.grade_id != kp.grade_id or parent.subject_id != kp.subject_id:
                    raise BusinessException("父级知识点必须与当前知识点同年级、同学科")
                if parent.id == kp.id or any(
                        d.id == parent.id for d in _subtree_nodes(db, kp)):
                    raise BusinessException("不能移动到自身或子节点下")
            if _sibling_exists(db, kp.grade_id, kp.subject_id, new_parent_id,
                               kp.name, exclude_id=kp.id):
                raise BusinessException("同级已存在同名知识点")
            kp.parent_id = parent.id if parent else None

    db.flush()
    if parent_changed:
        _refresh_level_path(db, kp)
        db.flush()
    db.commit()
    db.refresh(kp)
    index_knowledge(kp)
    return ok(_single_node(db, kp), message="知识点已更新")


@router.delete("/knowledge/{kp_id}", summary="删除知识点")
def delete_knowledge(
    kp_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    kp = db.get(KnowledgePoint, kp_id)
    if not kp:
        raise NotFoundException("知识点不存在")
    _assert_write(db, user, kp.grade_id, kp.subject_id)

    child_cnt = db.scalar(
        select(func.count()).select_from(KnowledgePoint).where(
            KnowledgePoint.parent_id == kp.id)
    ) or 0
    if child_cnt:
        raise BusinessException("该知识点还有子节点，请先删除子级")
    q_cnt = db.scalar(
        select(func.count()).select_from(Question).where(
            Question.knowledge_point_id == kp.id)
    ) or 0
    if q_cnt:
        raise BusinessException(
            f"该知识点已被 {q_cnt} 道题目引用，请先调整相关题目后再删除")
    m_cnt = db.scalar(
        select(func.count()).select_from(KnowledgeMastery).where(
            KnowledgeMastery.knowledge_point_id == kp.id)
    ) or 0
    if m_cnt:
        raise BusinessException(
            f"该知识点已有 {m_cnt} 条学生掌握度记录引用，删除会导致学情数据丢失")
    r_cnt = db.scalar(
        select(func.count()).select_from(ReviewPlan).where(
            ReviewPlan.knowledge_point_id == kp.id)
    ) or 0
    if r_cnt:
        raise BusinessException(
            f"该知识点已有 {r_cnt} 条复习计划引用，请先处理相关计划后再删除")
    db.delete(kp)
    db.commit()
    # RAG 生命周期：删除后清理向量索引（失败只告警，不阻断业务）
    delete_knowledge_index(kp_id)
    return ok(message="知识点已删除")


# ================================================================ 批量操作
def _batch_delete(
    db: Session, nodes: List[KnowledgePoint], done: List[int],
    skipped: List[Dict[str, Any]],
) -> None:
    ids = [n.id for n in nodes]
    # 子节点计数：包含未选中节点；同批父+子时先删子再删父
    child_counts = dict(db.execute(
        select(KnowledgePoint.parent_id, func.count()).where(
            KnowledgePoint.parent_id.in_(ids)
        ).group_by(KnowledgePoint.parent_id)
    ).all())
    q_counts = dict(db.execute(
        select(Question.knowledge_point_id, func.count()).where(
            Question.knowledge_point_id.in_(ids)
        ).group_by(Question.knowledge_point_id)
    ).all())
    m_counts = dict(db.execute(
        select(KnowledgeMastery.knowledge_point_id, func.count()).where(
            KnowledgeMastery.knowledge_point_id.in_(ids)
        ).group_by(KnowledgeMastery.knowledge_point_id)
    ).all())
    r_counts = dict(db.execute(
        select(ReviewPlan.knowledge_point_id, func.count()).where(
            ReviewPlan.knowledge_point_id.in_(ids)
        ).group_by(ReviewPlan.knowledge_point_id)
    ).all())
    selected = set(ids)
    for kp in sorted(nodes, key=lambda x: (x.level or 0, x.id), reverse=True):
        remain_children = int(child_counts.get(kp.id, 0) or 0)
        if remain_children > 0:
            skipped.append({"id": kp.id, "reason": "存在子节点，请先删除子级"})
            continue
        if q_counts.get(kp.id):
            skipped.append({"id": kp.id, "reason": f"被 {q_counts[kp.id]} 道题目引用"})
            continue
        if m_counts.get(kp.id):
            skipped.append({"id": kp.id, "reason": "已有学生掌握度记录引用"})
            continue
        if r_counts.get(kp.id):
            skipped.append({"id": kp.id, "reason": "已有复习计划引用"})
            continue
        db.delete(kp)
        done.append(kp.id)
        if kp.parent_id in selected:
            child_counts[kp.parent_id] = max(
                0, int(child_counts.get(kp.parent_id, 0) or 0) - 1)


def _batch_move(
    db: Session, user: User, nodes: List[KnowledgePoint],
    payload: Dict[str, Any], done: List[int], skipped: List[Dict[str, Any]],
) -> None:
    target_grade = _to_int(payload.get("grade_id"))
    target_subject = _to_int(payload.get("subject_id"))
    cross_scope = bool(target_grade or target_subject)
    if cross_scope and not _is_sys_admin(user):
        raise PermissionException("仅系统管理员可跨年级/学科移动知识点")
    if cross_scope:
        if not target_grade or not target_subject:
            raise BusinessException("跨年级/学科移动需同时提供 grade_id 与 subject_id")
        grade = db.get(Grade, target_grade)
        subject = db.get(Subject, target_subject)
        if not grade or not subject:
            raise BusinessException("目标年级或学科不存在")
        if subject.grade_id != target_grade:
            raise BusinessException("目标学科与年级不匹配")

    raw_parent = payload.get("parent_id")
    parent_supplied = "parent_id" in payload
    parent_id = _to_int(raw_parent) if parent_supplied else None
    parent: Optional[KnowledgePoint] = None
    if parent_id:
        parent = db.get(KnowledgePoint, parent_id)
        if not parent:
            raise BusinessException("父级知识点不存在")

    moved_descendants: Set[int] = set()
    for kp in sorted(nodes, key=lambda x: (x.level or 0, x.id)):
        if kp.id in moved_descendants:
            skipped.append({"id": kp.id, "reason": "祖先节点已在本批移动中"})
            continue
        grade_id = target_grade or kp.grade_id
        subject_id = target_subject or kp.subject_id
        if parent is not None:
            # 父节点自身也在移动列表里会被后续重算打乱，直接拒绝
            if parent.id in {n.id for n in nodes}:
                skipped.append({"id": kp.id, "reason": "父节点也在移动列表中"})
                continue
            if parent.grade_id != grade_id or parent.subject_id != subject_id:
                skipped.append({"id": kp.id, "reason": "父级与目标年级/学科不一致"})
                continue
        descendants = _subtree_nodes(db, kp)
        subtree_ids = {kp.id} | {d.id for d in descendants}
        if parent is not None and parent.id in subtree_ids:
            skipped.append({"id": kp.id, "reason": "不能移动到自身或子节点下"})
            continue
        if _sibling_exists(
                db, grade_id, subject_id, parent.id if parent else None, kp.name,
                exclude_id=kp.id):
            skipped.append({"id": kp.id, "reason": "目标位置存在同名同级节点"})
            continue

        kp.parent_id = parent.id if parent else None
        kp.grade_id = grade_id
        kp.subject_id = subject_id
        if parent:
            kp.level = (parent.level or 0) + 1
            kp.path = f"{parent.path or parent.id}/{kp.id}"
        else:
            kp.level = 1
            kp.path = str(kp.id)
        node_map: Dict[int, KnowledgePoint] = {kp.id: kp}
        for child in descendants:
            child.grade_id = grade_id
            child.subject_id = subject_id
            node_map[child.id] = child
        for child in descendants:
            p = node_map.get(child.parent_id)
            if p:
                child.level = (p.level or 0) + 1
                child.path = f"{p.path or p.id}/{child.id}"
        moved_descendants |= subtree_ids
        done.append(kp.id)


def _sync_batch_indexes(db: Session, action: str, done: List[int]) -> None:
    """批量写操作后的向量索引同步（RAG 生命周期）。

    - delete：删除对应向量；
    - publish/archive/update：重建这些节点的向量（archive 由
      `index_knowledge` 内部转为删除，避免语义检索命中已归档内容）；
    - move：节点及其整棵子树的 path/grade/subject 都变了，逐个重建。
    索引失败只告警，不影响业务返回。
    """
    if action == "delete":
        for kp_id in done:
            delete_knowledge_index(kp_id)
        return
    for kp_id in done:
        kp = db.get(KnowledgePoint, kp_id)
        if not kp:
            continue
        index_knowledge(kp)
        if action == "move":
            for child in _subtree_nodes(db, kp):
                index_knowledge(child)


@router.post("/knowledge/batch", summary="批量操作知识点")
def batch_knowledge(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    action = str(data.get("action") or "").strip().lower()
    if action not in {"delete", "publish", "archive", "update", "move"}:
        raise BusinessException("action 仅支持 delete / publish / archive / update / move")
    raw_ids = data.get("ids") or []
    if not isinstance(raw_ids, list) or not raw_ids:
        raise BusinessException("请选择要操作的知识点")
    ids: List[int] = []
    for item in raw_ids:
        value = _to_int(item)
        if value is None:
            raise BusinessException("ids 必须为整数列表")
        if value not in ids:
            ids.append(value)

    found = {kp.id: kp for kp in db.scalars(
        select(KnowledgePoint).where(KnowledgePoint.id.in_(ids))
    ).all()}
    nodes = [found[i] for i in ids if i in found]
    skipped: List[Dict[str, Any]] = [
        {"id": i, "reason": "不存在"} for i in ids if i not in found]
    if not nodes:
        return ok({"done": [], "skipped": skipped})

    # 契约：同批只要有一个越权节点，整体 403（防部分写入歧义）
    pred = _write_predicate(db, user)
    for kp in nodes:
        if not pred(kp.grade_id, kp.subject_id):
            raise PermissionException("所选知识点中存在无权维护的节点，已拒绝整批操作")

    payload = data.get("payload") or {}
    if not isinstance(payload, dict):
        raise BusinessException("payload 必须为对象")
    done: List[int] = []

    if action in ("publish", "archive"):
        target_status = "published" if action == "publish" else "archived"
        for kp in nodes:
            kp.status = target_status
            done.append(kp.id)
    elif action == "update":
        allowed = {"status", "difficulty", "keywords", "textbook_version", "term"}
        changes = {k: v for k, v in payload.items() if k in allowed}
        if not changes:
            raise BusinessException("update 需要 payload 至少包含一个可更新字段")
        if "status" in changes:
            changes["status"] = _validate_status(changes["status"])
        if "difficulty" in changes:
            changes["difficulty"] = _validate_difficulty(changes["difficulty"])
        if "term" in changes:
            changes["term"] = _to_int(changes["term"], 1) or 1
        for kp in nodes:
            for key, value in changes.items():
                if key in ("keywords", "textbook_version"):
                    value = _clean_str(value)
                setattr(kp, key, value)
            done.append(kp.id)
    elif action == "delete":
        _batch_delete(db, nodes, done, skipped)
    else:  # move
        _batch_move(db, user, nodes, payload, done, skipped)

    db.commit()
    _sync_batch_indexes(db, action, done)
    return ok({"done": done, "skipped": skipped}, message=f"批量操作完成 {len(done)} 条")


# ================================================================ 导入
def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX_MIME,
        headers={
            "Content-Disposition":
                f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().lstrip("\ufeff")


def _parse_import_file(content: bytes, filename: str) -> List[Tuple[int, Dict[str, str]]]:
    """解析 xlsx/csv，返回 [(Excel 行号, {表头: 值})]。"""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text: Optional[str] = None
        for encoding in ("utf-8-sig", "gbk", "utf-8"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise BusinessException("CSV 文件编码无法识别，请使用 UTF-8 或 GBK")
        rows_all = list(csv.reader(io.StringIO(text)))
        if not rows_all:
            raise BusinessException("文件内容为空")
        headers = [_cell_text(h) for h in rows_all[0]]
        data_rows = [
            (idx + 2, {headers[i]: _cell_text(cell) for i, cell in enumerate(row) if i < len(headers)})
            for idx, row in enumerate(rows_all[1:])
        ]
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001  非 xlsx 内容
            raise BusinessException("文件解析失败，请使用模板导出的 xlsx 或 csv")
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise BusinessException("文件内容为空")
        headers = [_cell_text(h) for h in header_row]
        data_rows = []
        for idx, row in enumerate(iterator, start=2):
            data_rows.append((
                idx,
                {headers[i]: _cell_text(cell)
                 for i, cell in enumerate(row) if i < len(headers)},
            ))
        wb.close()
    else:
        raise BusinessException("仅支持 .xlsx 或 .csv 文件")

    missing = [h for h in ("年级", "学科") if h not in headers]
    if missing:
        raise BusinessException(f"模板缺少必需列：{'、'.join(missing)}")
    return data_rows


def _import_row_hierarchy(row: Dict[str, str]) -> List[str]:
    """按「单元/章节/知识点」归并层级：空列自动折叠，保持相对层级。"""
    return [x for x in (
        row.get("单元", ""), row.get("章节", ""), row.get("知识点", ""),
    ) if x]


@router.post("/knowledge/import", summary="导入知识点")
async def import_knowledge(
    file: UploadFile = File(...), dry_run: bool = Form(False),
    db: Session = Depends(get_db),
    user: User = Depends(require_perm(WRITE_PERM)),
):
    content = await file.read()
    data_rows = _parse_import_file(content, file.filename or "")

    grades = {g.name: g for g in db.scalars(select(Grade)).all()}
    subjects = {
        (s.grade_id, s.name): s for s in db.scalars(select(Subject)).all()
    }
    predicate = _write_predicate(db, user)
    existing = {
        (k.grade_id, k.subject_id, k.parent_id, k.name): k
        for k in db.scalars(select(KnowledgePoint)).all()
    }
    created = updated = 0
    skipped: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    seen: Set[Tuple[int, int, Tuple[str, ...]]] = set()

    for row_no, row in data_rows:
        if not any((row.get(h) or "").strip() for h in IMPORT_HEADERS):
            continue

        grade_name = (row.get("年级") or "").strip()
        subject_name = (row.get("学科") or "").strip()
        if not grade_name or not subject_name:
            errors.append({"row": row_no, "message": "缺少年级或学科"})
            continue
        grade = grades.get(grade_name)
        if not grade:
            errors.append({"row": row_no, "message": f"年级「{grade_name}」不存在"})
            continue
        subject = subjects.get((grade.id, subject_name))
        if not subject:
            errors.append({
                "row": row_no,
                "message": f"年级「{grade_name}」下不存在学科「{subject_name}」",
            })
            continue
        if not predicate(grade.id, subject.id):
            errors.append({
                "row": row_no,
                "message": f"无权维护 {grade_name}/{subject_name} 的知识点",
            })
            continue

        chain = _import_row_hierarchy(row)
        if not chain:
            errors.append({"row": row_no, "message": "单元/章节/知识点至少填写一级"})
            continue
        diff = (row.get("难度") or "medium").strip().lower() or "medium"
        if diff not in VALID_DIFFICULTY:
            errors.append({
                "row": row_no, "message": f"难度「{diff}」非法（easy/medium/hard）"})
            continue
        try:
            term = int(float(row.get("学期") or 1))
        except (TypeError, ValueError):
            errors.append({"row": row_no, "message": "学期必须为数字"})
            continue
        textbook = _clean_str(row.get("教材版本"))
        keywords = _clean_str(row.get("关键词"))
        description = _clean_str(row.get("说明"))
        code = _clean_str(row.get("编码"))

        target_key = (grade.id, subject.id, tuple(chain))
        if target_key in seen:
            skipped.append({"row": row_no, "reason": "同级重名（文件内重复）"})
            continue
        seen.add(target_key)

        # 逐级解析：已存在的复用，缺失的登记待创建
        parent: Optional[KnowledgePoint] = None
        pending: List[str] = []
        leaf: Optional[KnowledgePoint] = None
        for depth, node_name in enumerate(chain, start=1):
            cache_key = (
                grade.id, subject.id,
                parent.id if parent else None, node_name,
            )
            node = existing.get(cache_key)
            if node is None:
                pending = chain[depth - 1:]
                break
            parent = node
            if depth == len(chain):
                leaf = node

        if leaf is not None and code and code != leaf.code:
            dup = db.scalar(select(KnowledgePoint.id).where(
                KnowledgePoint.grade_id == grade.id,
                KnowledgePoint.subject_id == subject.id,
                KnowledgePoint.code == code,
                KnowledgePoint.id != leaf.id,
            ).limit(1))
            if dup:
                errors.append({"row": row_no, "message": f"编码「{code}」已存在"})
                continue
        if leaf is None and code and _code_exists(db, grade.id, subject.id, code):
            errors.append({"row": row_no, "message": f"编码「{code}」已存在"})
            continue

        row_keys: List[Tuple[int, int, Optional[int], str]] = []
        created_before_row = created
        try:
            # 自动补建缺失的中间层（单元/章节）
            base_parent = parent
            for offset, node_name in enumerate(pending):
                depth = len(chain) - len(pending) + offset + 1  # 1-based 层级
                last = depth == len(chain)
                node = KnowledgePoint(
                    name=node_name, grade_id=grade.id, subject_id=subject.id,
                    parent_id=base_parent.id if base_parent else None,
                    level=depth, term=term,
                    difficulty=diff if last else "medium",
                    keywords=keywords if last else None,
                    description=description if last else None,
                    code=code if last else None,
                    textbook_version=textbook,
                    status="published", source="import", owner_id=user.id,
                    sort=0, is_active=True,
                )
                db.add(node)
                db.flush()
                node.path = (
                    f"{base_parent.path or base_parent.id}/{node.id}"
                    if base_parent else str(node.id)
                )
                cache_key = (
                    grade.id, subject.id,
                    base_parent.id if base_parent else None, node_name,
                )
                row_keys.append(cache_key)
                existing[cache_key] = node
                base_parent = node
                created += 1
                if last:
                    leaf = node
            if pending == [] and leaf is not None:
                # 已存在的叶子：按行更新非空字段（供重复导入修正信息）
                changed = False
                for field, value in (
                    ("textbook_version", textbook), ("term", term),
                    ("difficulty", diff), ("keywords", keywords),
                    ("description", description), ("code", code),
                ):
                    if value in (None, "") or getattr(leaf, field) == value:
                        continue
                    setattr(leaf, field, value)
                    changed = True
                if changed:
                    updated += 1
        except Exception as exc:  # noqa: BLE001  单行失败不拖垮整批
            # 真实导入：此前各行已 commit，rollback 只回滚本行；
            # dry_run：整体尚未 commit，rollback 会清空已规划的写入，
            # 因此记录错误后中止后续行（预览结果以已处理行为准）。
            db.rollback()
            for key in row_keys:
                existing.pop(key, None)
            created = created_before_row
            errors.append({"row": row_no, "message": f"写入失败：{exc}"})
            if dry_run:
                break
            continue

        if not dry_run:
            db.commit()
        else:
            db.flush()

    if dry_run:
        db.rollback()
    return ok({
        "dry_run": bool(dry_run),
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }, message=("导入校验完成（未写库）" if dry_run else "导入完成"))


# ================================================================ 向量重建
@router.post("/knowledge/reindex", summary="重建向量索引")
def reindex(
    db: Session = Depends(get_db),
    _: User = Depends(require_perm(WRITE_PERM)),
):
    """全量重建向量索引。

    契约保留「教师可调用」（E2E 既有行为）+ resource:knowledge:write 权限点。
    传入 wall-clock deadline，到点后 reindex_all 在批间自行收尾并返回
    timed_out 标记，避免长时间挂起拖死调用方。
    """
    t0 = time.time()
    budget = settings.REINDEX_TIMEOUT_SEC
    stats = reindex_all(db, deadline=time.monotonic() + budget)
    elapsed = int((time.time() - t0) * 1000)
    stats["elapsed_ms"] = elapsed

    if stats.get("timed_out"):
        return ok(stats, message=(
            f"向量索引重建超时（>{budget}s），已中断并保留已完成部分。"
            "若为首次运行，多因需要下载 embedding 模型，请检查网络后重试"
        ))
    slow = elapsed > settings.REINDEX_WARN_MS
    return ok(
        stats,
        message=("向量索引重建完成" if not slow
                 else f"向量索引重建耗时偏长({elapsed}ms)，建议分批执行"),
    )
