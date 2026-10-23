"""题库与试卷接口。"""
import io
import json
import math
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, Body, Depends, File, Form, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, noload

from app.agents.rag import delete_question_index, index_question
from app.api.deps import has_perm, require_perm
from app.core.database import get_db
from app.core.exception import BusinessException, NotFoundException, PermissionException
from app.models import (
    Grade, KnowledgePoint, Notification, Paper, PaperItem, PaperShare, Question,
    QuestionBankType, QuestionStatus, Subject, Task, User,
)
from app.utils.common import PageParams, dictify, ok

router = APIRouter(prefix="/question", tags=["题库与试卷"])

QUESTION_TYPE_LABEL = {
    "single": "单选题", "multiple": "多选题", "judge": "判断题", "blank": "填空题",
    "short": "简答题", "calc": "计算题", "material": "材料分析题", "composition": "作文题",
}
# 历史别名：早期接口与自动化脚本曾用 subjective 指代主观题，
# 而枚举里实际是 short。保留映射，避免存量数据与旧客户端创建的数据
# 在前端下拉框里匹配不到（显示成原始英文），也避免直接 422 卡住旧用法。
QUESTION_TYPE_ALIAS = {
    "subjective": "short",   # 主观题 -> 简答题
    "essay": "composition",  # 论述/作文
}
DIFFICULTY_LABEL = {"easy": "简单", "medium": "中等", "hard": "困难"}

# 判断题固定选项：历史数据/导入模板的答案是「正确/错误」，
# 但答题页提交的是选项 key（A/B）。建题/导题时补齐标准选项，
# 避免导入的判断题没有可点选项；判分侧由 app.utils.answer 做等价映射。
DEFAULT_JUDGE_OPTIONS = [{"key": "A", "text": "正确"}, {"key": "B", "text": "错误"}]


def normalize_question_type(raw: Any, *, strict: bool = True) -> str:
    """把外部传入的题型归一化为合法枚举值。

    strict=True 时非法值抛BusinessException（写接口用，脏数据不应入库）；
    strict=False 时非法值回落为 'single'（展示/兜底场景用）。
    """
    qtype = str(raw or "single").strip().lower()
    qtype = QUESTION_TYPE_ALIAS.get(qtype, qtype)
    if qtype in QUESTION_TYPE_LABEL:
        return qtype
    if strict:
        raise BusinessException(
            f"题型不合法：{raw!r}，可选值 {'/'.join(QUESTION_TYPE_LABEL)}"
        )
    return "single"


QUESTION_PERM_LIST = "question:bank:list"
QUESTION_PERM_WRITE = "question:bank:write"
QUESTION_PERM_IMPORT = "question:bank:import"
QUESTION_PERM_EXPORT = "question:bank:export"
QUESTION_PERM_AUDIT = "question:bank:audit"

QUESTION_STATUS_VALUES = {
    QuestionStatus.DRAFT.value, QuestionStatus.PUBLISHED.value,
    QuestionStatus.OFFLINE.value,
}
QUESTION_BANK_VALUES = {
    QuestionBankType.PRIVATE.value, QuestionBankType.PUBLIC.value,
}

SORT_MAP = {
    "created_desc": (Question.created_at.desc(), Question.id.desc()),
    "usage_desc": (Question.usage_count.desc(), Question.id.desc()),
    "updated_desc": (Question.updated_at.desc(), Question.id.desc()),
}


def _codes(user: User) -> set:
    return set(user.role_codes or [])


def _is_sys_admin(user: User) -> bool:
    return "sys_admin" in _codes(user)


def _is_grade_admin(user: User) -> bool:
    return "grade_admin" in _codes(user)


def _can_audit(user: User) -> bool:
    """公共题库上架/审核：sys_admin、grade_admin 或显式 question:bank:audit。"""
    if _is_sys_admin(user) or _is_grade_admin(user):
        return True
    return has_perm(user, QUESTION_PERM_AUDIT)


def _grade_scope(user: User) -> Optional[int]:
    """年级管理员的管辖年级；sys_admin 返回 None（不限）。"""
    if _is_sys_admin(user):
        return None
    return user.grade_admin_grade_id or user.grade_id or -1


def _visibility_clause(user: User):
    """题库可见性：sys_admin 全量；其他人「公共 + 本人私有」。

    关键修复：显式传 bank_type 时必须叠加本条件，否则传 bank_type=private
    会直接覆盖可见性过滤，任意教师都能拉到他人私有题。
    """
    if _is_sys_admin(user):
        return None
    return or_(
        Question.bank_type == QuestionBankType.PUBLIC.value,
        Question.owner_id == user.id,
    )


def _can_edit_question(q: Question, user: User) -> bool:
    """编辑/上下架/批量写：本人 / sys_admin / grade_admin（公共题审核场景）。"""
    if _is_sys_admin(user) or q.owner_id == user.id:
        return True
    if _is_grade_admin(user):
        return q.bank_type == QuestionBankType.PUBLIC.value
    return False


def _can_delete_question(q: Question, user: User) -> bool:
    """删除：本人 / sys_admin / grade_admin（公共题或本年级题目）。"""
    if _is_sys_admin(user) or q.owner_id == user.id:
        return True
    if _is_grade_admin(user):
        scope = _grade_scope(user)
        return q.bank_type == QuestionBankType.PUBLIC.value or (
            scope is not None and q.grade_id == scope)
    return False


# ---------------------------------------------------------------- 写接口校验
def _require_grade_subject(
    db: Session, grade_id: Any, subject_id: Any
) -> Tuple[Grade, Subject]:
    """create/update 必填校验：年级/学科存在且 subject.grade_id == grade_id。"""
    try:
        gid = int(grade_id)
        sid = int(subject_id)
    except (TypeError, ValueError):
        raise BusinessException("请选择年级和学科")
    if not gid or not sid:
        raise BusinessException("请选择年级和学科")
    grade = db.get(Grade, gid)
    if not grade:
        raise BusinessException("年级不存在")
    subject = db.get(Subject, sid)
    if not subject:
        raise BusinessException("学科不存在")
    if subject.grade_id != gid:
        raise BusinessException("学科与年级不匹配")
    return grade, subject


def _validate_difficulty(value: Any, default: str = "medium") -> str:
    diff = str(value or default).strip().lower()
    if diff not in DIFFICULTY_LABEL:
        raise BusinessException("难度只能是 easy / medium / hard")
    return diff


def _validate_score(value: Any, default: float = 5.0) -> float:
    try:
        score = float(value if value not in (None, "") else default)
    except (TypeError, ValueError):
        raise BusinessException("分值必须为数字")
    if not math.isfinite(score) or score <= 0 or score > 1000:
        raise BusinessException("分值需在 0~1000 之间")
    return score


def _validate_q_status(value: Any, default: str = "published") -> str:
    status = str(value or default).strip().lower()
    if status not in QUESTION_STATUS_VALUES:
        raise BusinessException("题目状态只能是 draft / published / offline")
    return status


def _validate_bank_type(value: Any, default: str = "private") -> str:
    bank = str(value or default).strip().lower()
    if bank not in QUESTION_BANK_VALUES:
        raise BusinessException("题库类型只能是 private / public")
    return bank


def _validate_knowledge_point(
    db: Session, kp_id: Any, grade_id: int, subject_id: int
) -> KnowledgePoint:
    """知识点必须存在，且与题目同年级同学科（防跨学科脏关联）。"""
    try:
        kid = int(kp_id)
    except (TypeError, ValueError):
        raise BusinessException("知识点不存在")
    kp = db.get(KnowledgePoint, kid)
    if not kp:
        raise BusinessException("知识点不存在")
    if kp.grade_id != grade_id or kp.subject_id != subject_id:
        raise BusinessException("知识点与题目的年级/学科不一致")
    return kp


def _filtered_question_stmt(
    user: User, *, keyword: Optional[str] = None, grade_id: Optional[int] = None,
    subject_id: Optional[int] = None, knowledge_point_id: Optional[int] = None,
    qtype: Optional[str] = None, difficulty: Optional[str] = None,
    bank_type: Optional[str] = None, status: Optional[str] = None,
    source: Optional[str] = None,
):
    """题库筛选（列表/统计/type_stats 共用同一套条件）。"""
    stmt = select(Question)
    visibility = _visibility_clause(user)
    if visibility is not None:
        stmt = stmt.where(visibility)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(Question.stem.like(like), Question.analysis.like(like)))
    if grade_id:
        stmt = stmt.where(Question.grade_id == grade_id)
    if subject_id:
        stmt = stmt.where(Question.subject_id == subject_id)
    if knowledge_point_id:
        stmt = stmt.where(Question.knowledge_point_id == knowledge_point_id)
    if qtype:
        stmt = stmt.where(Question.type == qtype)
    if difficulty:
        stmt = stmt.where(Question.difficulty == difficulty)
    if bank_type:
        stmt = stmt.where(Question.bank_type == bank_type)
    if status:
        stmt = stmt.where(Question.status == status)
    if source:
        stmt = stmt.where(Question.source == source)
    return stmt


def _question_out(
    q: Question, kp_names: Optional[Dict[int, str]] = None
) -> Dict[str, Any]:
    """题目出参：保留旧字段并补 type_label/difficulty_label/knowledge_point_name。"""
    d = dictify(q)
    d["type"] = normalize_question_type(q.type, strict=False)
    d["type_label"] = QUESTION_TYPE_LABEL.get(d["type"], d["type"])
    d["difficulty_label"] = DIFFICULTY_LABEL.get(q.difficulty, q.difficulty)
    d["knowledge_point_name"] = (kp_names or {}).get(q.knowledge_point_id)
    return d


def _kp_name_map(db: Session, rows: List[Question]) -> Dict[int, str]:
    """一次预取知识点名称，避免列表循环内逐题 db.get 的 N+1。"""
    kp_ids = {q.knowledge_point_id for q in rows if q.knowledge_point_id}
    if not kp_ids:
        return {}
    return {
        kp.id: kp.name for kp in db.scalars(
            select(KnowledgePoint).where(KnowledgePoint.id.in_(kp_ids))
        ).all()
    }


# ---------------------------------------------------------------- 题目
@router.get("/questions", summary="题目列表")
def list_questions(
    params: PageParams = Depends(),
    keyword: Optional[str] = None,
    grade_id: Optional[int] = None,
    subject_id: Optional[int] = None,
    knowledge_point_id: Optional[int] = None,
    qtype: Optional[str] = None,
    difficulty: Optional[str] = None,
    bank_type: Optional[str] = None,
    status: Optional[str] = None,
    source: Optional[str] = None,
    sort: str = "created_desc",
    db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_LIST)),
):
    stmt = _filtered_question_stmt(
        user, keyword=keyword, grade_id=grade_id, subject_id=subject_id,
        knowledge_point_id=knowledge_point_id, qtype=qtype,
        difficulty=difficulty, bank_type=bank_type, status=status, source=source,
    )
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = db.scalars(
        stmt.order_by(*SORT_MAP.get(sort, SORT_MAP["created_desc"]))
        .offset(params.offset).limit(params.page_size)
    ).all()

    kp_names = _kp_name_map(db, rows)
    items = [_question_out(q, kp_names) for q in rows]

    # type_stats 必须与当前筛选一致（排除 qtype 维度本身，便于切换题型标签）
    type_base = _filtered_question_stmt(
        user, keyword=keyword, grade_id=grade_id, subject_id=subject_id,
        knowledge_point_id=knowledge_point_id, qtype=None,
        difficulty=difficulty, bank_type=bank_type, status=status, source=source,
    )
    type_stats = {
        normalize_question_type(t, strict=False): int(c or 0)
        for t, c in db.execute(
            type_base.with_only_columns(Question.type, func.count())
            .group_by(Question.type)
        ).all()
    }

    return ok({
        "items": items, "total": total, "page": params.page, "page_size": params.page_size,
        "type_stats": type_stats,
    })


@router.get("/questions/stats", summary="题库统计")
def question_stats(
    keyword: Optional[str] = None,
    grade_id: Optional[int] = None,
    subject_id: Optional[int] = None,
    knowledge_point_id: Optional[int] = None,
    qtype: Optional[str] = None,
    difficulty: Optional[str] = None,
    bank_type: Optional[str] = None,
    status: Optional[str] = None,
    source: Optional[str] = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_LIST)),
):
    """与列表同筛选范围（不分页）的题库统计。"""
    stmt = _filtered_question_stmt(
        user, keyword=keyword, grade_id=grade_id, subject_id=subject_id,
        knowledge_point_id=knowledge_point_id, qtype=qtype,
        difficulty=difficulty, bank_type=bank_type, status=status, source=source,
    )

    def _count(extra=None) -> int:
        target = stmt if extra is None else stmt.where(extra)
        return db.scalar(select(func.count()).select_from(target.subquery())) or 0

    def _group(column) -> Dict[str, int]:
        rows = db.execute(
            stmt.with_only_columns(column, func.count()).group_by(column)
        ).all()
        return {str(key): int(cnt or 0) for key, cnt in rows}

    return ok({
        "total": _count(),
        "by_type": _group(Question.type),
        "by_difficulty": _group(Question.difficulty),
        "by_status": _group(Question.status),
        "by_bank": _group(Question.bank_type),
        "with_analysis": _count(Question.analysis.isnot(None) & (Question.analysis != "")),
        "with_knowledge": _count(Question.knowledge_point_id.isnot(None)),
        "ai_source": _count(Question.source.like("ai%")),
    })


@router.post("/questions", summary="新增题目")
def create_question(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_WRITE)),
):
    stem = str(data.get("stem") or "").strip()
    if not stem:
        raise BusinessException("题干不能为空")
    grade, subject = _require_grade_subject(db, data.get("grade_id"), data.get("subject_id"))
    qtype = normalize_question_type(data.get("type", "single"))
    difficulty = _validate_difficulty(data.get("difficulty"))
    score = _validate_score(data.get("score"))
    status = _validate_q_status(data.get("status"))
    bank_type = _validate_bank_type(data.get("bank_type"))
    kp_id = None
    if data.get("knowledge_point_id") not in (None, ""):
        kp = _validate_knowledge_point(
            db, data["knowledge_point_id"], grade.id, subject.id)
        kp_id = kp.id
    # 公共题库：非审核角色创建/发布一律进待审核（draft），修复「教师可自行上架公共题」
    if bank_type == QuestionBankType.PUBLIC.value and status == "published" \
            and not _can_audit(user):
        status = QuestionStatus.DRAFT.value

    options = data.get("options")
    if qtype == "judge" and not options:
        options = DEFAULT_JUDGE_OPTIONS
    q = Question(
        type=qtype,
        stem=stem,
        options=_dump(options), answer=_dump(data.get("answer")),
        analysis=data.get("analysis"), score=score,
        difficulty=difficulty,
        grade_id=grade.id, subject_id=subject.id,
        knowledge_point_id=kp_id,
        bank_type=bank_type,
        owner_id=user.id, status=status,
        fingerprint=_fingerprint(stem), source="manual",
    )
    db.add(q)
    db.commit()
    db.refresh(q)
    index_question(q)
    return ok(_question_out(q, _kp_name_map(db, [q])), message="题目创建成功")


@router.post("/questions/batch", summary="批量操作题目")
def batch_questions(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_WRITE)),
):
    """批量删除/上下架/移动题库/分配知识点/修改难度。

    权限：同批只要有一道题无权操作，整体 403（防部分写入歧义）；
    业务失败（被试卷引用、知识点不匹配等）进 skipped，不影响其它题。
    """
    action = str(data.get("action") or "").strip().lower()
    if action not in {"delete", "publish", "offline", "draft", "move_bank",
                      "assign_knowledge", "set_difficulty"}:
        raise BusinessException(
            "action 仅支持 delete / publish / offline / draft / move_bank / "
            "assign_knowledge / set_difficulty")
    raw_ids = data.get("ids") or []
    if not isinstance(raw_ids, list) or not raw_ids:
        raise BusinessException("请选择要操作的题目")
    ids: List[int] = []
    for item in raw_ids:
        try:
            value = int(item)
        except (TypeError, ValueError):
            raise BusinessException("ids 必须为整数列表")
        if value not in ids:
            ids.append(value)

    found = {
        q.id: q for q in db.scalars(
            select(Question).where(Question.id.in_(ids))).all()
    }
    nodes = [found[i] for i in ids if i in found]
    skipped: List[Dict[str, Any]] = [
        {"id": i, "reason": "不存在"} for i in ids if i not in found]
    if not nodes:
        return ok({"done": [], "skipped": skipped})

    permission_check = (
        _can_delete_question if action == "delete" else _can_edit_question)
    if any(not permission_check(q, user) for q in nodes):
        raise PermissionException("所选题目中存在无权操作的题目，已拒绝整批操作")

    payload = data.get("payload") or {}
    if not isinstance(payload, dict):
        raise BusinessException("payload 必须为对象")
    done: List[int] = []
    rag_delete_ids: List[int] = []

    if action == "delete":
        ref_counts = dict(db.execute(
            select(PaperItem.question_id, func.count()).where(
                PaperItem.question_id.in_([q.id for q in nodes])
            ).group_by(PaperItem.question_id)
        ).all())
        for q in nodes:
            ref = int(ref_counts.get(q.id, 0) or 0)
            if ref:
                skipped.append({
                    "id": q.id,
                    "reason": f"已被 {ref} 份试卷引用，请先从相关试卷中移除",
                })
                continue
            db.delete(q)
            done.append(q.id)
            rag_delete_ids.append(q.id)
    elif action in ("publish", "offline", "draft"):
        target = {
            "publish": QuestionStatus.PUBLISHED.value,
            "offline": QuestionStatus.OFFLINE.value,
            "draft": QuestionStatus.DRAFT.value,
        }[action]
        for q in nodes:
            if target == QuestionStatus.PUBLISHED.value \
                    and q.bank_type == QuestionBankType.PUBLIC.value \
                    and not _can_audit(user):
                skipped.append({"id": q.id, "reason": "公共题库上架需管理员审核"})
                continue
            q.status = target
            done.append(q.id)
    elif action == "move_bank":
        if payload.get("bank_type") in (None, ""):
            raise BusinessException("move_bank 需要 payload.bank_type（private/public）")
        bank = _validate_bank_type(payload.get("bank_type"))
        for q in nodes:
            q.bank_type = bank
            # 移入公共题库：非审核角色强制待审核，审核角色保留当前状态
            if bank == QuestionBankType.PUBLIC.value and not _can_audit(user):
                q.status = QuestionStatus.DRAFT.value
            done.append(q.id)
    elif action == "assign_knowledge":
        if "knowledge_point_id" not in payload:
            raise BusinessException("assign_knowledge 需要 payload.knowledge_point_id")
        raw_kp = payload.get("knowledge_point_id")
        kp = None
        if raw_kp not in (None, ""):
            try:
                kp = db.get(KnowledgePoint, int(raw_kp))
            except (TypeError, ValueError):
                kp = None
            if not kp:
                raise BusinessException("知识点不存在")
        for q in nodes:
            if kp is not None and (
                    kp.grade_id != q.grade_id or kp.subject_id != q.subject_id):
                skipped.append({
                    "id": q.id, "reason": "知识点与题目的年级/学科不一致"})
                continue
            q.knowledge_point_id = kp.id if kp else None
            done.append(q.id)
    else:  # set_difficulty
        if payload.get("difficulty") in (None, ""):
            raise BusinessException("set_difficulty 需要 payload.difficulty")
        difficulty = _validate_difficulty(payload.get("difficulty"))
        for q in nodes:
            q.difficulty = difficulty
            done.append(q.id)

    db.commit()
    # RAG 生命周期：删除成功后清理对应题目向量（失败仅告警，不阻断）
    for qid in rag_delete_ids:
        delete_question_index(qid)
    return ok({"done": done, "skipped": skipped},
              message=f"批量操作完成 {len(done)} 道，跳过 {len(skipped)} 道")


@router.get("/questions/template", summary="下载批量导入模板")
def import_template(
    db: Session = Depends(get_db),
    _: User = Depends(require_perm(QUESTION_PERM_IMPORT)),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "题目导入"
    headers = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "答案", "解析",
               "难度", "分值", "知识点ID", "题库类型"]
    ws.append(headers)
    ws.append(["single", "1 + 1 = ?", "1", "2", "3", "4", "B",
               "加法运算", "easy", "5", "", "private"])
    ws.append(["judge", "地球是圆的", "", "", "", "", "正确", "常识", "easy", "5", "", "private"])
    ws.append(["short", "简述光合作用的意义", "", "", "", "", "", "", "medium", "10", "", "private"])
    for i, w in enumerate(headers, start=1):
        ws.column_dimensions[chr(64 + i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=questions_template.xlsx"},
    )


@router.put("/questions/{q_id}", summary="编辑题目")
def update_question(
    q_id: int, data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_WRITE)),
):
    q = db.get(Question, q_id)
    if not q:
        raise NotFoundException("题目不存在")
    if not _can_edit_question(q, user):
        raise PermissionException("只能编辑自己的题目")

    # 年级/学科不允许单独改到不匹配组合：先算出生效值并整体校验
    effective_grade = q.grade_id
    effective_subject = q.subject_id
    if "grade_id" in data:
        try:
            effective_grade = int(data["grade_id"])
        except (TypeError, ValueError):
            raise BusinessException("请选择年级")
        if not db.get(Grade, effective_grade):
            raise BusinessException("年级不存在")
    if "subject_id" in data:
        try:
            effective_subject = int(data["subject_id"])
        except (TypeError, ValueError):
            raise BusinessException("请选择学科")
        if not db.get(Subject, effective_subject):
            raise BusinessException("学科不存在")
    if "grade_id" in data or "subject_id" in data:
        _require_grade_subject(db, effective_grade, effective_subject)

    if "stem" in data:
        stem = str(data.get("stem") or "").strip()
        if not stem:
            raise BusinessException("题干不能为空")
        q.stem = stem
        q.fingerprint = _fingerprint(stem)
    if "type" in data and data.get("type") is not None:
        q.type = normalize_question_type(data["type"])
    if "difficulty" in data:
        q.difficulty = _validate_difficulty(data["difficulty"], default=q.difficulty)
    if "score" in data:
        q.score = _validate_score(data["score"], default=q.score)
    if "knowledge_point_id" in data:
        raw_kp = data.get("knowledge_point_id")
        if raw_kp in (None, ""):
            q.knowledge_point_id = None
        else:
            q.knowledge_point_id = _validate_knowledge_point(
                db, raw_kp, effective_grade, effective_subject).id
    if "grade_id" in data:
        q.grade_id = effective_grade
    if "subject_id" in data:
        q.subject_id = effective_subject
    if "options" in data:
        q.options = _dump(data.get("options"))
    if "answer" in data:
        q.answer = _dump(data.get("answer"))
    if "analysis" in data:
        q.analysis = data.get("analysis")
    # status 不在此处更新：上下架/审核必须走 /status 端点，
    # 否则编辑接口会绕过「公共题库上架需管理员审核」的校验（E2E 依赖此行为）
    # 题型切到判断题但没传选项时补齐标准选项，保证答题页可作答
    if q.type == "judge" and not q.options:
        q.options = _dump(DEFAULT_JUDGE_OPTIONS)
    db.commit()
    db.refresh(q)
    index_question(q)
    return ok(_question_out(q, _kp_name_map(db, [q])), message="题目已更新")


@router.delete("/questions/{q_id}", summary="删除题目")
def delete_question(
    q_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_WRITE)),
):
    q = db.get(Question, q_id)
    if not q:
        raise NotFoundException("题目不存在")
    if not _can_delete_question(q, user):
        raise PermissionException("只能删除自己的题目")
    # 引用检查：已组卷的题目删除会让 PaperItem 悬空、试卷总分虚高、
    # 考试预览缺题——必须先从相关试卷移除
    ref_cnt = db.scalar(select(func.count()).select_from(PaperItem).where(
        PaperItem.question_id == q_id)) or 0
    if ref_cnt:
        raise BusinessException(f"该题目已被 {ref_cnt} 份试卷引用，请先从相关试卷中移除")
    db.delete(q)
    db.commit()
    # RAG 生命周期：删除后清理题目向量（失败只告警，不阻断业务）
    delete_question_index(q_id)
    return ok(message="题目已删除")


@router.put("/questions/{q_id}/status", summary="上下架 / 审核")
def change_status(
    q_id: int, data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_WRITE)),
):
    q = db.get(Question, q_id)
    if not q:
        raise NotFoundException("题目不存在")
    if not _can_edit_question(q, user):
        raise PermissionException("只能操作自己的题目")
    status = _validate_q_status(data.get("status"), default="")
    if status == QuestionStatus.PUBLISHED.value \
            and q.bank_type == QuestionBankType.PUBLIC.value \
            and not _can_audit(user):
        # 公共题库上架需管理员审核权限
        raise PermissionException("公共题库上架需管理员审核")
    q.status = status
    db.commit()
    return ok(message="状态已更新")


@router.get("/questions/{q_id}/duplicates", summary="题目查重")
def check_duplicate(
    q_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_LIST)),
):
    q = db.get(Question, q_id)
    if not q:
        raise NotFoundException("题目不存在")
    visibility = _visibility_clause(user)
    if visibility is not None and not (
            q.bank_type == QuestionBankType.PUBLIC.value or q.owner_id == user.id):
        raise PermissionException("无权查看该题目")
    stmt = select(Question).where(
        Question.fingerprint == q.fingerprint, Question.id != q.id)
    if visibility is not None:
        stmt = stmt.where(visibility)
    dupes = db.scalars(stmt.limit(10)).all()
    return ok([{"id": d.id, "stem": d.stem[:80], "created_at": d.created_at} for d in dupes])


@router.post("/questions/import", summary="Excel 批量导入题目")
async def import_questions(
    request: Request,
    file: UploadFile = File(...),
    grade_id: Optional[int] = Form(None),
    subject_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_IMPORT)),
):
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise BusinessException(f"Excel 解析失败：{exc}")

    # 前后端契约：前端以 multipart 表单字段传年级/学科（FormData.append）；
    # 同时兼容旧脚本用 query string 传参。
    if grade_id is None:
        try:
            grade_id = int(request.query_params.get("grade_id") or 0) or None
        except (TypeError, ValueError):
            grade_id = None
    if subject_id is None:
        try:
            subject_id = int(request.query_params.get("subject_id") or 0) or None
        except (TypeError, ValueError):
            subject_id = None

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise BusinessException("文件为空")

    headers = [str(h).strip() if h else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(headers)}
    required = {"题型", "题干", "答案"}
    if not required.issubset(set(headers)):
        raise BusinessException(f"模板缺少必需列：{'、'.join(required - set(headers))}")

    # 导入必须落在明确的年级/学科下（缺 grade 不再默认 0）；只给学科时按学科反推年级
    if not subject_id:
        raise BusinessException("请选择导入的学科")
    subject = db.get(Subject, int(subject_id))
    if not subject:
        raise BusinessException("学科不存在")
    if grade_id and int(grade_id) != subject.grade_id:
        raise BusinessException("学科与年级不匹配")
    target_grade = subject.grade_id
    target_subject = subject.id

    created, skipped = 0, []
    seen_fps: set = set()
    for line_no, row in enumerate(rows[1:], start=2):
        def get(key: str, default=None):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else default

        stem = get("题干")
        if not stem or not str(stem).strip():
            continue
        qtype = str(get("题型", "single")).strip()
        if qtype not in QUESTION_TYPE_LABEL:
            qtype = "single"
        options = None
        opts = {k: get(f"选项{k}") for k in "ABCD"}
        if any(opts.values()):
            options = [{"key": k, "text": str(v)} for k, v in opts.items() if v]
            options = json.dumps(options, ensure_ascii=False)
        # 导入模板的判断题没有选项列：补齐 A/B，否则答题页无选项可点
        if qtype == "judge" and not options:
            options = json.dumps(DEFAULT_JUDGE_OPTIONS, ensure_ascii=False)
        difficulty = str(get("难度", "medium") or "medium").strip().lower()
        if difficulty not in DIFFICULTY_LABEL:
            skipped.append({"line": line_no, "reason": f"难度非法：{difficulty}",
                            "stem": str(stem)[:40]})
            continue
        try:
            score = _validate_score(get("分值"), default=5.0)
        except BusinessException as exc:
            skipped.append({"line": line_no, "reason": exc.message,
                            "stem": str(stem)[:40]})
            continue
        kp_id = None
        raw_kp = get("知识点ID")
        if raw_kp not in (None, ""):
            try:
                kp = _validate_knowledge_point(db, raw_kp, target_grade, target_subject)
                kp_id = kp.id
            except BusinessException as exc:
                skipped.append({"line": line_no, "reason": exc.message,
                                "stem": str(stem)[:40]})
                continue
        bank_type = str(get("题库类型", "private") or "private").strip().lower()
        if bank_type not in QUESTION_BANK_VALUES:
            bank_type = QuestionBankType.PRIVATE.value
        fp = _fingerprint(str(stem))
        if fp in seen_fps or db.scalar(
                select(Question).where(Question.fingerprint == fp)):
            skipped.append({"line": line_no, "reason": "题目重复", "stem": str(stem)[:40]})
            continue
        seen_fps.add(fp)
        db.add(Question(
            type=qtype, stem=str(stem), options=options,
            answer=_dump(get("答案")), analysis=get("解析"),
            difficulty=difficulty, score=score,
            grade_id=target_grade, subject_id=target_subject,
            knowledge_point_id=kp_id,
            bank_type=bank_type,
            # 公共题库导入必须走审核（draft），防止教师批量导入直接上架；
            # 私有题库导入即上架（仅自己可见），与单题创建行为一致
            owner_id=user.id,
            status=(QuestionStatus.DRAFT.value
                    if bank_type == QuestionBankType.PUBLIC.value else
                    QuestionStatus.PUBLISHED.value),
            fingerprint=fp, source="import",
        ))
        created += 1
    db.commit()
    return ok({"created": created, "skipped": skipped, "skipped_count": len(skipped)},
              message=f"成功导入 {created} 道题目，跳过 {len(skipped)} 道")


@router.get("/questions/export", summary="导出题目为 Excel")
def export_questions(
    subject_id: Optional[int] = None, grade_id: Optional[int] = None,
    keyword: Optional[str] = None, difficulty: Optional[str] = None,
    bank_type: Optional[str] = None, status: Optional[str] = None,
    source: Optional[str] = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_perm(QUESTION_PERM_EXPORT)),
):
    """导出题目（沿用列表筛选与可见性；最多 5000 行）。"""
    stmt = _filtered_question_stmt(
        user, keyword=keyword, grade_id=grade_id, subject_id=subject_id,
        difficulty=difficulty, bank_type=bank_type, status=status, source=source,
    )
    rows = db.scalars(stmt.limit(5000)).all()

    # 一次预取知识点名称，避免循环内逐题 db.get(KnowledgePoint) 的 N+1
    kp_ids = {q.knowledge_point_id for q in rows if q.knowledge_point_id}
    kp_names = {
        kp.id: kp.name for kp in db.scalars(
            select(KnowledgePoint).where(KnowledgePoint.id.in_(kp_ids))
        ).all()
    } if kp_ids else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "题库导出"
    ws.append(["题型", "题干", "答案", "解析", "难度", "分值", "知识点ID", "状态", "创建时间"])
    for q in rows:
        ws.append([
            QUESTION_TYPE_LABEL.get(q.type, q.type), q.stem, q.answer, q.analysis,
            DIFFICULTY_LABEL.get(q.difficulty, q.difficulty), q.score,
            kp_names.get(q.knowledge_point_id, "") if q.knowledge_point_id else "",
            q.status,
            q.created_at.strftime("%Y-%m-%d %H:%M") if q.created_at else "",
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=questions_export.xlsx"},
    )


# ---------------------------------------------------------------- 试卷
def _clean_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _dump_json_field(value: Any) -> Optional[str]:
    """difficulty_ratio 等 String(64) 字段：dict/list 统一转 JSON 字符串。"""
    if value in (None, ""):
        return None
    if isinstance(value, str):
        return value[:64]
    return json.dumps(value, ensure_ascii=False)[:64]


def _paper_visibility_clause(user: User):
    """试卷可见性：sys_admin 全量；其他用户 = 本人创建 / 公开共享 / PaperShare 接收。"""
    if _is_sys_admin(user):
        return None
    return or_(
        Paper.creator_id == user.id,
        Paper.is_shared.is_(True),
        Paper.id.in_(select(PaperShare.paper_id).where(
            PaperShare.to_user_id == user.id)),
    )


def _can_edit_paper(paper: Paper, user: User) -> bool:
    """编辑/删除：创建者 / sys_admin；grade_admin 仅公共卷（is_shared）或本人。"""
    if _is_sys_admin(user) or paper.creator_id == user.id:
        return True
    if _is_grade_admin(user):
        return bool(paper.is_shared)
    return False


def _validate_paper_name(value: Any) -> str:
    name = str(value or "").strip()
    if not name:
        raise BusinessException("试卷名称不能为空")
    if len(name) > 128:
        raise BusinessException("试卷名称不能超过 128 个字符")
    return name


def _validate_duration(value: Any, default: int = 60) -> int:
    try:
        duration = int(value if value not in (None, "") else default)
    except (TypeError, ValueError):
        raise BusinessException("考试时长必须为整数分钟")
    if duration < 1 or duration > 300:
        raise BusinessException("考试时长需在 1~300 分钟之间")
    return duration


def _normalize_paper_items(
    db: Session, user: User, raw_items: Any, grade_id: int, subject_id: int,
) -> Tuple[List[Dict[str, Any]], float]:
    """组卷/编辑题目校验：存在、可访问、同年级同学科、不重复、分值合法。

    返回 (规范化 items, 总分)。总分一律由 items 求和，忽略客户端 total_score。
    """
    if not isinstance(raw_items, list) or not raw_items:
        raise BusinessException("请至少选择一道题目")
    seen: set = set()
    prepared: List[Dict[str, Any]] = []
    total = 0.0
    for raw in raw_items:
        if not isinstance(raw, dict):
            raise BusinessException("题目数据格式错误")
        try:
            qid = int(raw.get("question_id"))
        except (TypeError, ValueError):
            raise BusinessException("题目数据格式错误：question_id 无效")
        if qid in seen:
            raise BusinessException(f"试卷中题目重复（question_id={qid}）")
        seen.add(qid)
        q = db.get(Question, qid)
        if not q:
            raise BusinessException(f"题目不存在或已删除（id={qid}）")
        # 可访问题目：公共题库或本人私有（共享试卷不扩大题库权限）
        if not _is_sys_admin(user) and not (
                q.bank_type == QuestionBankType.PUBLIC.value
                or q.owner_id == user.id):
            raise BusinessException(f"无权使用题目（id={qid}）：仅可选公共题或本人私有题")
        if q.grade_id != grade_id or q.subject_id != subject_id:
            raise BusinessException(f"题目年级/学科与试卷不一致（id={qid}）")
        try:
            score = float(raw.get("score")
                          if raw.get("score") is not None else (q.score or 5))
        except (TypeError, ValueError):
            raise BusinessException(f"题目分值必须为数字（id={qid}）")
        if not math.isfinite(score) or score < 0 or score > 1000:
            raise BusinessException(f"题目分值需在 0~1000 之间（id={qid}）")
        section = raw.get("section_name")
        prepared.append({
            "question_id": qid,
            "score": round(score, 2),
            "section_name": str(section)[:64] if section else None,
        })
        total += score
    return prepared, round(total, 2)


def _paper_item_rows(db: Session, paper_ids: List[int]) -> Dict[int, int]:
    """分组统计题量（不含题分，供列表/统计复用）。"""
    if not paper_ids:
        return {}
    rows = db.execute(
        select(PaperItem.paper_id, func.count(PaperItem.id))
        .where(PaperItem.paper_id.in_(paper_ids))
        .group_by(PaperItem.paper_id)
    ).all()
    return {int(pid): int(cnt or 0) for pid, cnt in rows}


@router.get("/papers", summary="试卷列表")
def list_papers(
    params: PageParams = Depends(),
    subject_id: Optional[int] = None,
    grade_id: Optional[int] = None,
    keyword: Optional[str] = None,
    source: Optional[str] = None,
    shared: Optional[bool] = None,
    creator_id: Optional[int] = None,
    sort: str = "created_desc",
    db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:list")),
):
    stmt = _paper_filtered_stmt(
        user, subject_id=subject_id, grade_id=grade_id, keyword=keyword,
        source=source, shared=shared, creator_id=creator_id,
    )
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    order = (
        (Paper.updated_at.desc(), Paper.id.desc())
        if sort == "updated_desc"
        else (Paper.created_at.desc(), Paper.id.desc())
    )
    rows = db.scalars(
        # noload：列表只需要题量，不再连带加载整卷题目明细
        stmt.options(noload(Paper.items))
        .order_by(*order).offset(params.offset).limit(params.page_size)
    ).all()
    paper_ids = [p.id for p in rows]
    item_counts = _paper_item_rows(db, paper_ids)

    # 批量预取创建者 / 学科 / 年级名称，避免逐卷 db.get 的 N+1
    creator_ids = {p.creator_id for p in rows if p.creator_id}
    subject_ids = {p.subject_id for p in rows if p.subject_id}
    grade_ids = {p.grade_id for p in rows if p.grade_id}
    creator_names = {
        u.id: u.real_name for u in db.scalars(
            select(User).where(User.id.in_(creator_ids))
        ).all()
    } if creator_ids else {}
    subject_names = {
        s.id: s.name for s in db.scalars(
            select(Subject).where(Subject.id.in_(subject_ids))
        ).all()
    } if subject_ids else {}
    grade_names = {
        g.id: g.name for g in db.scalars(
            select(Grade).where(Grade.id.in_(grade_ids))
        ).all()
    } if grade_ids else {}

    items = []
    for p in rows:
        d = dictify(p)
        d["item_count"] = item_counts.get(p.id, 0)
        d["creator_name"] = creator_names.get(p.creator_id)
        d["subject_name"] = subject_names.get(p.subject_id)
        d["grade_name"] = grade_names.get(p.grade_id)
        # AI 卷标记：来源以 ai 开头（ai / ai_import 等）
        d["ai_source"] = (p.source or "").startswith("ai")
        items.append(d)
    return ok({"items": items, "total": total, "page": params.page})


def _paper_filtered_stmt(
    user: User, *, subject_id: Optional[int] = None,
    grade_id: Optional[int] = None, keyword: Optional[str] = None,
    source: Optional[str] = None, shared: Optional[bool] = None,
    creator_id: Optional[int] = None,
):
    """试卷列表/统计共用筛选条件（含可见性）。"""
    stmt = select(Paper)
    visibility = _paper_visibility_clause(user)
    if visibility is not None:
        stmt = stmt.where(visibility)
    if subject_id:
        stmt = stmt.where(Paper.subject_id == subject_id)
    if grade_id:
        stmt = stmt.where(Paper.grade_id == grade_id)
    if keyword:
        stmt = stmt.where(Paper.name.like(f"%{keyword}%"))
    if source:
        stmt = stmt.where(Paper.source == source)
    if shared is not None:
        stmt = stmt.where(Paper.is_shared.is_(bool(shared)))
    # creator_id 仅管理员可用（普通用户可见性本就限定为本人/共享，忽略该参数）
    if creator_id and _is_sys_admin(user):
        stmt = stmt.where(Paper.creator_id == creator_id)
    return stmt


@router.get("/papers/stats", summary="试卷统计")
def paper_stats(
    subject_id: Optional[int] = None, grade_id: Optional[int] = None,
    keyword: Optional[str] = None, source: Optional[str] = None,
    shared: Optional[bool] = None, creator_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:list")),
):
    """同筛选范围（不分页）的试卷统计：总数/题量/平均分/分布/共享/AI 卷。"""
    stmt = _paper_filtered_stmt(
        user, subject_id=subject_id, grade_id=grade_id, keyword=keyword,
        source=source, shared=shared, creator_id=creator_id,
    )
    rows = db.execute(
        stmt.with_only_columns(
            Paper.id, Paper.subject_id, Paper.grade_id, Paper.source,
            Paper.is_shared, Paper.total_score,
        )
    ).all()
    total = len(rows)
    paper_ids = [r[0] for r in rows]
    item_counts = _paper_item_rows(db, paper_ids)
    total_items = sum(item_counts.values())
    scores = [float(r[5] or 0) for r in rows]
    avg_score = round(sum(scores) / total, 2) if total else 0

    subject_ids = {r[1] for r in rows if r[1]}
    grade_ids = {r[2] for r in rows if r[2]}
    subject_names = {
        s.id: s.name for s in db.scalars(
            select(Subject).where(Subject.id.in_(subject_ids))).all()
    } if subject_ids else {}
    grade_names = {
        g.id: g.name for g in db.scalars(
            select(Grade).where(Grade.id.in_(grade_ids))).all()
    } if grade_ids else {}

    by_subject: Dict[str, int] = {}
    by_grade: Dict[str, int] = {}
    by_source: Dict[str, int] = {}
    shared_count = 0
    ai_count = 0
    for r in rows:
        subject_name = subject_names.get(r[1]) or "未设置"
        grade_name = grade_names.get(r[2]) or "未设置"
        src = r[3] or "manual"
        by_subject[subject_name] = by_subject.get(subject_name, 0) + 1
        by_grade[grade_name] = by_grade.get(grade_name, 0) + 1
        by_source[src] = by_source.get(src, 0) + 1
        if r[4]:
            shared_count += 1
        if src.startswith("ai"):
            ai_count += 1

    return ok({
        "total": total,
        "total_items": total_items,
        "avg_score": avg_score,
        "by_subject": by_subject,
        "by_grade": by_grade,
        "by_source": by_source,
        "shared": shared_count,
        "ai_papers": ai_count,
    })


@router.post("/papers/batch", summary="批量操作试卷")
def batch_papers(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:write")),
):
    """批量删除/复制/共享试卷。返回 {done, skipped}；同批无权整体 403。"""
    action = str(data.get("action") or "").strip().lower()
    if action not in {"delete", "copy", "share"}:
        raise BusinessException("action 仅支持 delete / copy / share")
    raw_ids = data.get("ids") or []
    if not isinstance(raw_ids, list) or not raw_ids:
        raise BusinessException("请选择要操作的试卷")
    ids: List[int] = []
    for item in raw_ids:
        try:
            value = int(item)
        except (TypeError, ValueError):
            raise BusinessException("ids 必须为整数列表")
        if value not in ids:
            ids.append(value)

    found = {
        p.id: p for p in db.scalars(
            select(Paper).where(Paper.id.in_(ids))).all()
    }
    papers = [found[i] for i in ids if i in found]
    skipped: List[Dict[str, Any]] = [
        {"id": i, "reason": "不存在"} for i in ids if i not in found]
    if not papers:
        return ok({"done": [], "skipped": skipped})

    if action == "share" and not has_perm(user, "paper:share"):
        raise PermissionException("缺少权限：paper:share")

    # 同批无权限整体 403：delete/copy 按编辑/可读判定；share 按契约非 owner 进 skipped
    if action == "delete":
        if any(not _can_edit_paper(p, user) for p in papers):
            raise PermissionException("所选试卷中存在无权删除的试卷，已拒绝整批操作")
    elif action == "copy":
        for p in papers:
            try:
                _assert_paper_readable(db, p, user)
            except PermissionException:
                raise PermissionException("所选试卷中存在无权复制的试卷，已拒绝整批操作")

    payload = data.get("payload") or {}
    if not isinstance(payload, dict):
        raise BusinessException("payload 必须为对象")
    done: List[int] = []
    copied: List[Dict[str, Any]] = []

    if action == "delete":
        ref_counts = dict(db.execute(
            select(Task.paper_id, func.count()).where(
                Task.paper_id.in_([p.id for p in papers])
            ).group_by(Task.paper_id)
        ).all())
        for p in papers:
            ref = int(ref_counts.get(p.id, 0) or 0)
            if ref:
                skipped.append({
                    "id": p.id,
                    "reason": f"已被 {ref} 个作业/考试引用，请先处理相关任务",
                })
                continue
            db.delete(p)
            done.append(p.id)
    elif action == "copy":
        for p in papers:
            new = _copy_paper_row(db, p, user)
            done.append(p.id)
            copied.append({"source_id": p.id, "id": new.id, "name": new.name})
    else:  # share
        user_ids = payload.get("user_ids") or (
            [payload["to_user_id"]] if payload.get("to_user_id") else [])
        targets = _validate_share_targets(db, user_ids)
        for p in papers:
            if not (_is_sys_admin(user) or p.creator_id == user.id):
                skipped.append({"id": p.id, "reason": "仅创建者或管理员可共享"})
                continue
            _share_paper_to(db, p, user, targets)
            done.append(p.id)

    db.commit()
    result: Dict[str, Any] = {"done": done, "skipped": skipped}
    if copied:
        result["copied"] = copied
    return ok(result, message=f"批量操作完成 {len(done)} 份，跳过 {len(skipped)} 份")


@router.get("/papers/{paper_id}", summary="试卷详情（含题目）")
def get_paper(
    paper_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:list")),
):
    paper = db.get(Paper, paper_id)
    if not paper:
        raise NotFoundException("试卷不存在")
    _assert_paper_readable(db, paper, user)

    d = dictify(paper)
    items = []
    for it in sorted(paper.items, key=lambda x: x.sort_order):
        q = it.question
        kp = db.get(KnowledgePoint, q.knowledge_point_id) if q and q.knowledge_point_id else None
        qtype = normalize_question_type(q.type, strict=False) if q else None
        items.append({
            "paper_item_id": it.id, "question_id": q.id, "sort_order": it.sort_order,
            "score": it.score, "section_name": it.section_name,
            "type": qtype, "type_label": QUESTION_TYPE_LABEL.get(qtype, qtype),
            "stem": q.stem, "options": _load(q.options), "answer": _load(q.answer),
            "analysis": q.analysis, "difficulty": q.difficulty,
            "knowledge_point_name": kp.name if kp else None,
        })
    d["items"] = items
    # 补充展示字段（旧字段全部保留）
    d["creator_name"] = db.scalar(
        select(User.real_name).where(User.id == paper.creator_id))
    d["subject_name"] = db.scalar(
        select(Subject.name).where(Subject.id == paper.subject_id))
    d["grade_name"] = db.scalar(
        select(Grade.name).where(Grade.id == paper.grade_id))
    d["item_count"] = len(items)
    d["ai_source"] = (paper.source or "").startswith("ai")
    return ok(d)


@router.post("/papers", summary="手动组卷")
def create_paper(
    data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:write")),
):
    name = _validate_paper_name(data.get("name"))
    grade, subject = _require_grade_subject(
        db, data.get("grade_id"), data.get("subject_id"))
    duration = _validate_duration(data.get("duration"))
    items, total = _normalize_paper_items(
        db, user, data.get("items"), grade.id, subject.id)
    # 总分由 items 求和，忽略客户端伪造的 total_score；source 固定 manual
    # （AI 组卷由 paper_agent 落库时写 source=ai，不走本接口）
    paper = Paper(
        name=name, grade_id=grade.id, subject_id=subject.id,
        total_score=total, duration=duration, creator_id=user.id,
        source="manual", difficulty_ratio=_dump_json_field(data.get("difficulty_ratio")),
    )
    db.add(paper)
    db.flush()
    for i, item in enumerate(items, start=1):
        db.add(PaperItem(
            paper_id=paper.id, question_id=item["question_id"],
            sort_order=i, score=item["score"], section_name=item["section_name"],
        ))
    db.commit()
    db.refresh(paper)
    return ok({"id": paper.id, "name": paper.name, "total_score": paper.total_score},
              message="组卷成功")


@router.put("/papers/{paper_id}", summary="编辑试卷")
def update_paper(
    paper_id: int, data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:write")),
):
    paper = db.get(Paper, paper_id)
    if not paper:
        raise NotFoundException("试卷不存在")
    if not _can_edit_paper(paper, user):
        raise PermissionException("只能编辑自己创建的试卷")

    if "name" in data:
        paper.name = _validate_paper_name(data.get("name"))
    effective_grade = paper.grade_id
    effective_subject = paper.subject_id
    if "grade_id" in data or "subject_id" in data:
        try:
            if "grade_id" in data:
                effective_grade = int(data["grade_id"])
            if "subject_id" in data:
                effective_subject = int(data["subject_id"])
        except (TypeError, ValueError):
            raise BusinessException("请选择年级和学科")
        grade, subject = _require_grade_subject(db, effective_grade, effective_subject)
        effective_grade, effective_subject = grade.id, subject.id
    if "duration" in data:
        paper.duration = _validate_duration(data.get("duration"), default=paper.duration)

    if "items" in data:
        items, total = _normalize_paper_items(
            db, user, data.get("items"), effective_grade, effective_subject)
        db.execute(PaperItem.__table__.delete().where(
            PaperItem.paper_id == paper_id))
        for i, item in enumerate(items, start=1):
            db.add(PaperItem(
                paper_id=paper_id, question_id=item["question_id"],
                sort_order=i, score=item["score"], section_name=item["section_name"],
            ))
        paper.total_score = total
    elif "grade_id" in data or "subject_id" in data:
        # 仅改年级/学科而题目不变时，必须保证现存题目与新年级/学科一致
        mismatched = db.scalar(
            select(func.count()).select_from(PaperItem)
            .join(Question, Question.id == PaperItem.question_id)
            .where(
                PaperItem.paper_id == paper_id,
                or_(Question.grade_id != effective_grade,
                    Question.subject_id != effective_subject),
            )
        ) or 0
        if mismatched:
            raise BusinessException(
                f"试卷中仍有 {mismatched} 道题目的年级/学科与新设置不一致，请先调整题目")

    paper.grade_id = effective_grade
    paper.subject_id = effective_subject
    if "difficulty_ratio" in data:
        paper.difficulty_ratio = _dump_json_field(data.get("difficulty_ratio"))
    # 不改变 creator_id / source（契约要求）
    db.commit()
    return ok({"id": paper.id, "total_score": paper.total_score}, message="试卷已更新")


@router.delete("/papers/{paper_id}", summary="删除试卷")
def delete_paper(
    paper_id: int, db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:write")),
):
    paper = db.get(Paper, paper_id)
    if not paper:
        raise NotFoundException("试卷不存在")
    if not _can_edit_paper(paper, user):
        raise PermissionException("只能删除自己创建的试卷")
    # 外键 edu_task.paper_id 是 ON DELETE CASCADE：直接删会连带物理删除
    # 已发布的作业/考试及全部学生提交与成绩（实测 1 份试卷可带走 14 份提交），
    # 因此必须先阻断，由用户显式处理任务后再删。
    task_count = db.scalar(
        select(func.count()).select_from(Task).where(Task.paper_id == paper_id)
    ) or 0
    if task_count:
        raise BusinessException(
            f"该试卷已被 {task_count} 个作业/考试引用，删除会连同学生作答与成绩一起丢失；"
            "请先删除相关任务或改为停用"
        )
    db.delete(paper)
    db.commit()
    return ok(message="试卷已删除")


def _copy_paper_row(db: Session, src: Paper, user: User, name: Optional[str] = None) -> Paper:
    new = Paper(
        name=(name or f"{src.name}副本")[:128], grade_id=src.grade_id,
        subject_id=src.subject_id, total_score=src.total_score,
        duration=src.duration, difficulty_ratio=src.difficulty_ratio,
        knowledge_coverage=src.knowledge_coverage, difficulty_report=src.difficulty_report,
        creator_id=user.id, source=src.source,
    )
    db.add(new)
    db.flush()
    total = 0.0
    for it in src.items:
        db.add(PaperItem(
            paper_id=new.id, question_id=it.question_id, sort_order=it.sort_order,
            score=it.score, section_name=it.section_name,
        ))
        total += float(it.score or 0)
    new.total_score = round(total, 2)
    return new


@router.post("/papers/{paper_id}/copy", summary="试卷复制复用")
def copy_paper(
    paper_id: int, data: Dict[str, Any] = Body(default={}),
    db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:write")),
):
    src = db.get(Paper, paper_id)
    if not src:
        raise NotFoundException("试卷不存在")
    _assert_paper_readable(db, src, user)
    new = _copy_paper_row(db, src, user, name=_clean_str(data.get("name")))
    db.commit()
    return ok({"id": new.id, "name": new.name}, message="试卷复制成功")


def _validate_share_targets(db: Session, user_ids: Any) -> List[int]:
    """分享目标校验：去重、必须是教师/管理员账号。"""
    if not isinstance(user_ids, list) or not user_ids:
        raise BusinessException("请选择接收人")
    targets: List[int] = []
    for raw in user_ids:
        try:
            uid = int(raw)
        except (TypeError, ValueError):
            raise BusinessException("接收人参数无效")
        if uid in targets:
            continue
        user = db.get(User, uid)
        if not user:
            raise BusinessException(f"接收人不存在（id={uid}）")
        if user.identity not in ("teacher", "grade_admin", "sys_admin"):
            raise BusinessException(f"接收人必须是教师（id={uid}）")
        targets.append(uid)
    if not targets:
        raise BusinessException("请选择接收人")
    return targets


def _share_paper_to(db: Session, paper: Paper, user: User, targets: List[int]) -> None:
    """共享试卷并发送站内通知（重复共享仍通知，分享记录去重）。"""
    for uid in targets:
        exists = db.scalar(
            select(PaperShare.id).where(
                PaperShare.paper_id == paper.id, PaperShare.to_user_id == uid))
        if not exists:
            db.add(PaperShare(
                paper_id=paper.id, from_user_id=user.id, to_user_id=uid))
        db.add(Notification(
            user_id=uid, title="收到共享试卷",
            content=f"{user.real_name} 共享了试卷《{paper.name}》",
            type="system", biz_type="paper", biz_id=paper.id,
        ))
    paper.is_shared = True


@router.post("/papers/{paper_id}/share", summary="共享试卷给同科教师")
def share_paper(
    paper_id: int, data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:share")),
):
    paper = db.get(Paper, paper_id)
    if not paper:
        raise NotFoundException("试卷不存在")
    if paper.creator_id != user.id and not _is_sys_admin(user):
        raise PermissionException("只能共享自己的试卷")
    # 兼容旧契约：user_ids 列表 或 to_user_id 单值
    raw_ids = data.get("user_ids")
    if not raw_ids and data.get("to_user_id"):
        raw_ids = [data.get("to_user_id")]
    targets = _validate_share_targets(db, raw_ids or [])
    _share_paper_to(db, paper, user, targets)
    db.commit()
    return ok(message=f"已共享给 {len(targets)} 名教师")


@router.get("/papers/{paper_id}/export", summary="导出试卷（Excel/Word/PDF）")
def export_paper(
    paper_id: int, fmt: str = Query("excel", regex="^(excel|word|pdf|html)$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_perm("paper:export")),
):
    """导出试卷：
    - excel：xlsx 文件（默认）
    - word/pdf/html：HTML 文档（浏览器打印为 PDF，Word 可直接打开为 .doc）
    轻量化方案：HTML 兼容 Word 与 PDF 打印，无需额外依赖。
    """
    paper = db.get(Paper, paper_id)
    if not paper:
        raise NotFoundException("试卷不存在")
    _assert_paper_readable(db, paper, user)

    if fmt == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "试卷"
        sub = db.get(Subject, paper.subject_id)
        gr = db.get(Grade, paper.grade_id)
        ws.append([paper.name])
        ws.append([f"{gr.name if gr else ''} {sub.name if sub else ''}",
                   f"总分：{paper.total_score}", f"时间：{paper.duration} 分钟"])
        ws.append([])
        current_section = None
        idx = 1
        for it in sorted(paper.items, key=lambda x: x.sort_order):
            if it.section_name and it.section_name != current_section:
                current_section = it.section_name
                ws.append([current_section])
            q = it.question
            opts = _load(q.options) or []
            opt_text = "  ".join(f"{o['key']}. {o['text']}" for o in opts)
            ws.append([f"{idx}. （{it.score}分）{q.stem}"])
            if opt_text:
                ws.append([f"    {opt_text}"])
            ws.append([f"    答案：{q.answer or ''}"])
            if q.analysis:
                ws.append([f"    解析：{q.analysis}"])
            ws.append([])
            idx += 1
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=paper_{paper_id}.xlsx"},
        )

    # word / pdf / html → HTML 文档
    sub = db.get(Subject, paper.subject_id)
    gr = db.get(Grade, paper.grade_id)
    rows = ["<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<title>{paper.name}</title></head><body>"]
    rows.append(f"<h1 style='text-align:center'>{paper.name}</h1>")
    rows.append(f"<p style='text-align:center'>{gr.name if gr else ''} {sub.name if sub else ''}"
                f"　总分：{paper.total_score}　时间：{paper.duration} 分钟</p><hr/>")
    current_section = None
    idx = 1
    for it in sorted(paper.items, key=lambda x: x.sort_order):
        if it.section_name and it.section_name != current_section:
            current_section = it.section_name
            rows.append(f"<h3>{current_section}</h3>")
        q = it.question
        opts = _load(q.options) or []
        rows.append(f"<p><b>{idx}.</b>（{it.score}分）{q.stem}</p>")
        if opts:
            rows.append("<ol type='A'>" + "".join(
                f"<li>{o.get('text', '')}</li>" for o in opts) + "</ol>")
        rows.append(f"<p><i>答案：{q.answer or ''}</i></p>")
        if q.analysis:
            rows.append(f"<p><i>解析：{q.analysis}</i></p><hr/>")
        idx += 1
    rows.append("</body></html>")
    html = "\n".join(rows)
    ext = "doc" if fmt == "word" else "html"
    mt = "application/msword" if fmt == "word" else "text/html"
    return StreamingResponse(
        io.BytesIO(html.encode("utf-8")), media_type=mt,
        headers={"Content-Disposition": f"attachment; filename=paper_{paper_id}.{ext}"},
    )


# ---------------------------------------------------------------- 工具
def _assert_paper_readable(db: Session, paper: Paper, user: User) -> None:
    """试卷读取归属校验：仅创建者、被共享者、公开卷或管理员可读。

    试卷详情含标准答案，缺少校验会导致任意教师遍历 paper_id
    读取他人试卷答案。复制与导出同样受此保护。
    """
    if paper.creator_id == user.id or "sys_admin" in user.role_codes:
        return
    if paper.is_shared:
        return
    shared = db.scalar(
        select(PaperShare.id).where(
            PaperShare.paper_id == paper.id,
            PaperShare.to_user_id == user.id,
        )
    )
    if not shared:
        raise PermissionException("无权查看该试卷")


def _dump(v: Any) -> Optional[str]:
    if v is None:
        return None
    return v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)


def _load(v: Optional[str]) -> Any:
    if not v:
        return None
    try:
        return json.loads(v)
    except json.JSONDecodeError:
        return v


def _fingerprint(stem: str) -> str:
    import hashlib
    import re

    norm = re.sub(r"\s+", "", re.sub(r"[（(].*?[)）]", "", stem))
    return hashlib.md5(norm.encode()).hexdigest()
